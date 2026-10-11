"""交付质量闸（ADR-065 M4）：产出后按中枢判据做三层自检。

**判据来源（硬约束）**：三层闸的任务判据一律只来自中枢下发的 ``specification``
（``checklist`` 优先，其次 ``spec`` / ``tasks`` 原文），节点**不得自编任务判据**
（不得凭空要求某些章节、关键词、数据点或引用数量）。本模块只做三件事：把中枢
判据解析成可执行检查、执行检查、把未通过条目结构化输出（供回灌与阻断上报）。
唯一的引擎级判据是「禁占位符 / 禁空章节」——它由 ADR-065 M4 对 L1 直接规定，
不随任务变化，也不改变任何任务要求。

分层（对齐中枢分层验收口径；任一层不过**立即短路**并返回首个未通过层）：

- **L0 结构闸（硬闸）**：交付物存在 → 能被对应库**真打开**（python-docx /
  python-pptx / openpyxl，按扩展名）→ 内容块数不低于非空下限；
- **L1 内容闸（硬闸）**：checklist 可判定项自动抽取 —— 必需章节/关键词命中率、
  必需数据点（正则/数值）、引用 URL 数 ≥ N（可选 HEAD 200 可达性）、禁占位符、禁空章节；
- **L2 语义闸**：LLM 按 checklist 剩余条目**逐条判**并**必须给出证据位置**
  （第几段/第几张幻灯片）；只给结论不给证据位置一律视为不通过。L2 判据只来自
  checklist（无 checklist 时不做语义判分，避免把 spec 正文的过程描述当验收要求）。

阈值说明：``DEFAULT_MIN_BLOCKS``（非空下限）与 ``PLACEHOLDER_PATTERNS`` 是引擎下限，
不是新增任务判据；中枢 checklist 可用「至少 N 段/张/行」覆盖非空下限。

回灌纪律（ADR-065 M4）：未通过时只把**未通过条目 + 期望 + 实际**作为新一轮输入
（``VerifyReport.repair_instruction``），已达标内容不得重写；≤ ``MAX_REPAIR_ROUNDS``
轮仍不过 → 由调用方上报 ``agent.task_verify_blocked`` 阻断（不再发起下一轮）。
"""

import asyncio
import json
import logging
import math
import re
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import httpx

from app.executor.file_scanner import is_excluded_relative_path, scan_file_tree
from app.executor.llm_client import LLMError
from app.tools.base import ToolError
from app.tools.read_tool import resolve_read_path

logger = logging.getLogger(__name__)

#: 分层标识（封闭枚举，与中枢分层验收报告口径一致）
LAYER_STRUCTURE = "L0"
LAYER_CONTENT = "L1"
LAYER_SEMANTIC = "L2"
#: 全部层通过时的层标识（报告 layer 字段取值之一）
LAYER_PASSED = "passed"

#: 回灌修复轮上限（ADR-065 M4：≤5 轮；用尽仍不过 → verify_blocked）
MAX_REPAIR_ROUNDS = 5

#: 结构闸「非空」下限（引擎下限，中枢 checklist 的「至少 N 段」可覆盖）
DEFAULT_MIN_BLOCKS = 1

#: 可结构校验的 Office 交付物：扩展名 → (解析库名, 内容块单位)
OFFICE_READERS: dict[str, tuple[str, str]] = {
    ".docx": ("python-docx", "段落/表格"),
    ".pptx": ("python-pptx", "幻灯片"),
    ".xlsx": ("openpyxl", "数据行"),
}
#: 纯文本交付物（按行计内容块）
TEXT_SUFFIXES: tuple[str, ...] = (".md", ".markdown", ".txt", ".csv", ".tsv", ".json")
#: 交付物扫描候选扩展名
ARTIFACT_SUFFIXES: tuple[str, ...] = (*OFFICE_READERS, *TEXT_SUFFIXES)
#: 交付物点名（checklist/spec 中带引号的文件名）允许的扩展名
_NAMED_SUFFIXES = (*ARTIFACT_SUFFIXES, ".png", ".jpg", ".jpeg", ".pdf")

#: 禁占位符模式（ADR-065 M4 L1 硬闸；标签 → 正则）
PLACEHOLDER_PATTERNS: tuple[tuple[str, str], ...] = (
    ("TODO", r"\bTODO\b"),
    ("TBD", r"\bTBD\b"),
    ("undefined", r"\bundefined\b"),
    ("{{…}} 模板占位", r"\{\{[^{}\n]{0,80}\}\}"),
    ("XXX 占位", r"\bXXX+\b"),
    ("待补充/待填写", r"待补充|待填写|待完善|此处省略"),
    ("尖括号占位", r"<[^<>\n]{0,20}(?:待补充|待填写|占位|补充)[^<>\n]{0,20}>"),
)

#: L2 判分提示中交付物正文的总字符上限（防撑爆判分模型上下文）
_L2_MAX_CONTENT_CHARS = 16_000
#: 上报字段截断长度（requirement/reason 等）
_MAX_FIELD_CHARS = 300
#: 产物 mtime 与任务起始时刻的比较容差（秒，规避文件系统时间粒度）
_MTIME_TOLERANCE = 1.0
#: URL 可达性探测超时（秒）
_URL_PROBE_TIMEOUT = 8.0


# --------------------------------------------------------------------------- #
# 数据结构
# --------------------------------------------------------------------------- #


@dataclass
class VerifyItem:
    """单条判据结果（回灌与上报的最小单元：{id, layer, requirement, passed, evidence|reason}）。"""

    id: str
    layer: str
    requirement: str
    passed: bool
    evidence: str = ""
    reason: str = ""

    def to_dict(self) -> dict[str, Any]:
        """转结构化字典（通过给 ``evidence``，未通过给 ``reason``）。"""
        return {
            "id": self.id,
            "layer": self.layer,
            "requirement": self.requirement,
            "passed": self.passed,
            "evidence": _clip(self.evidence),
            "reason": _clip(self.reason),
        }


@dataclass
class ChecklistItem:
    """一条 checklist 原文条目（判据唯一来源）。"""

    id: str
    text: str


@dataclass
class ContentRule:
    """L1 可自动判定的内容判据（由 checklist 条目解析而来）。"""

    kind: str  # keywords | datapoint | urls | placeholder
    item_id: str  # 来源 checklist 条目 ID（如 C3）
    requirement: str  # checklist 原文（回灌与上报展示用）
    keywords: list[str] = field(default_factory=list)
    min_hits: int = 0
    pattern: str = ""
    expected: str = ""
    literal: bool = True
    min_count: int = 0
    #: checklist 明确要求「链接可访问」时才做 HEAD 探测（需调用方显式开启）
    check_url_status: bool = False


@dataclass
class VerifyPlan:
    """三层判据计划：只由中枢下发的规格/checklist 解析，节点不自编任务判据。"""

    source: str = "none"  # checklist | spec | none
    checklist: list[ChecklistItem] = field(default_factory=list)
    content_rules: list[ContentRule] = field(default_factory=list)
    semantic: list[ChecklistItem] = field(default_factory=list)
    #: 中枢点名的交付物（引号内文件名；空表示按本轮产物扫描）
    deliverables: list[str] = field(default_factory=list)
    #: 引擎默认非空下限（未写该类下限时生效）
    min_blocks: int = DEFAULT_MIN_BLOCKS
    #: 按内容块单位细分的非空下限（「至少 3 行」只作用于数据行类交付物）
    min_blocks_by_unit: dict[str, int] = field(default_factory=dict)

    def min_blocks_for(self, unit: str) -> int:
        """该交付物单位的非空下限（checklist 未写该类下限时用引擎默认）。"""
        return self.min_blocks_by_unit.get(_UNIT_KEYS.get(unit, ""), self.min_blocks)

    @property
    def has_criteria(self) -> bool:
        """是否存在中枢下发的判据（无判据时闸门不生效）。"""
        return bool(self.checklist or self.deliverables)

    @classmethod
    def from_specification(cls, specification: Any | None) -> "VerifyPlan":
        """从中枢下发的规格三件套解析判据（checklist 优先，其次 spec/tasks）。"""
        checklist = _field(specification, "checklist")
        spec_text = _field(specification, "spec")
        tasks_text = _field(specification, "tasks")
        source_text = (checklist or "").strip() or (spec_text or "").strip() or (
            tasks_text or ""
        ).strip()
        source = "checklist" if (checklist or "").strip() else ("spec" if source_text else "none")
        plan = cls.from_text(source_text, source=source)
        # L2 语义闸按 checklist 逐条判（ADR-065 M4）：无 checklist 时不做语义判分，
        # 避免把 spec/tasks 正文里的过程性描述当验收要求逐条判（L0/L1 判据仍生效）
        if source != "checklist":
            plan.semantic = []
        # 交付物点名同时参考 spec/tasks（checklist 常只写验收口径）
        extra = "\n".join(part for part in (spec_text, tasks_text) if part)
        for name in _parse_deliverables(extra):
            if name not in plan.deliverables:
                plan.deliverables.append(name)
        return plan

    @classmethod
    def from_text(cls, text: str | None, *, source: str = "checklist") -> "VerifyPlan":
        """从判据原文解析计划（测试与直接调用入口）。"""
        texts = parse_checklist_items(text)
        items = [ChecklistItem(f"C{index}", value) for index, value in enumerate(texts, 1)]
        rules, semantic = _parse_content_rules(items)
        return cls(
            source=source if items else "none",
            checklist=items,
            content_rules=rules,
            semantic=semantic,
            deliverables=_parse_deliverables(text or ""),
            min_blocks_by_unit=_parse_min_blocks(text or ""),
        )


@dataclass
class VerifyReport:
    """三层闸报告：``passed`` + 首个未通过层 ``layer`` + 逐条 ``items``。"""

    passed: bool
    layer: str
    items: list[VerifyItem] = field(default_factory=list)
    source: str = "none"
    #: 已回灌修复轮数（0 表示首次判定）
    repair_round: int = 0
    #: 附加说明（如闸门未生效、判分输出不可解析）
    note: str = ""

    @property
    def failures(self) -> list[VerifyItem]:
        """未通过条目（回灌输入与阻断上报的唯一来源）。"""
        return [item for item in self.items if not item.passed]

    def to_dict(self) -> dict[str, Any]:
        """规格要求的结构化报告：{passed, layer, items:[...]}。"""
        return {
            "passed": self.passed,
            "layer": self.layer,
            "items": [item.to_dict() for item in self.items],
        }

    def summary(self) -> str:
        """人类可读摘要（回灌提示、上报 reason 与前端展示用）。"""
        if self.passed:
            return "交付质量闸通过（L0 结构 / L1 内容 / L2 语义）"
        lines = [
            f"交付质量闸未通过：卡在 {self.layer}（已回灌修复 {self.repair_round} 轮）；"
            f"未通过 {len(self.failures)} 条："
        ]
        lines.extend(
            f"❌ [{item.layer}/{item.id}] {item.requirement} —— 实际：{item.reason or '未达标'}"
            for item in self.failures
        )
        return "\n".join(lines)

    def repair_instruction(self, round_index: int, max_rounds: int = MAX_REPAIR_ROUNDS) -> str:
        """构造回灌指令：**只含未通过条目**（条目 + 期望 + 实际），已达标项不得重写。"""
        blocks = [
            f"【{item.layer} / {item.id}】{item.requirement}\n"
            f"期望：{item.requirement}\n"
            f"实际：{item.reason or '未达标'}"
            for item in self.failures
        ]
        joined = "\n\n".join(blocks) or "（无具体未通过条目）"
        return (
            "自动交付质量闸未通过，请只修复下列未通过条目（自动流程，无需询问用户）：\n"
            f"（第 {round_index}/{max_rounds} 轮）\n\n{joined}\n\n"
            "要求：\n"
            "1. **只修改上面列出的未通过条目**涉及的内容；已达标内容必须原样保留，"
            "不得重写、改写、精简或删除；\n"
            "2. 修复后重新调用写工具覆盖同一交付物文件（保持原文件名）；\n"
            "3. 不得改动验收要求本身来「绕过」未通过项；\n"
            "4. 若某项确实无法达标，请在最终说明中写明原因，不要伪造通过。"
        )

    def to_blocked_report(self) -> dict[str, Any]:
        """构造上报/落库用报告（字段对齐中枢 ``verify_report`` 既有结构）。

        ``must_pass_failures`` 每条含前端展示所需 ``layer/name/cmd/status`` 与
        「卡在哪一条」明细 ``id/expected/actual``。
        """
        return {
            "status": "passed" if self.passed else "failed",
            "source": self.source,
            "layer": self.layer,
            "must_pass_failures": [
                {
                    "layer": item.layer,
                    "name": _clip(item.requirement),
                    "cmd": _clip(item.reason or item.evidence),
                    "status": "failed",
                    "id": item.id,
                    "expected": _clip(item.requirement),
                    "actual": _clip(item.reason),
                }
                for item in self.failures
            ],
            "checks": [item.to_dict() for item in self.items],
            "repair_rounds": self.repair_round,
            "summary": self.summary(),
        }

    def to_blocked_payload(self, task_id: str, session_id: str = "") -> dict[str, Any]:
        """构造 ``agent.task_verify_blocked`` 的完整 WS payload。

        字段对齐中枢既有契约（``TaskVerifyBlockedPayload`` 与
        ``_handle_task_verify_blocked``）：``task_id`` / ``session_id`` /
        ``reason`` / ``report``。
        """
        payload: dict[str, Any] = {
            "task_id": task_id,
            "reason": self.summary()[:2000],
            "report": self.to_blocked_report(),
        }
        if session_id:
            payload["session_id"] = session_id
        return payload


# --------------------------------------------------------------------------- #
# 判据解析（checklist → VerifyPlan）
# --------------------------------------------------------------------------- #

#: checklist 条目行（- / * / + / 1. / 1) / - [ ] 形式）
_ITEM_LINE_RE = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s*(?:\[[ xX]\]\s*)?(?P<text>.+?)\s*$")
#: 表格行（| a | b |）
_TABLE_ROW_RE = re.compile(r"^\s*\|(?P<cells>.+)\|\s*$")
#: 行内标记清洗（加粗/代码标记）
_MARKUP_RE = re.compile(r"[*`]+")
#: 引号内的文件名（中文无词边界，只认引号点名，避免把散文误当文件名）
_QUOTED_NAME_RE = re.compile(
    r"[「『“\"'`《](?P<name>[^」』”\"'`》\n]{1,120}?\.(?:"
    + "|".join(suffix.lstrip(".") for suffix in _NAMED_SUFFIXES)
    + r"))[」』”\"'`》]",
    re.IGNORECASE,
)
#: 非空下限（「至少 5 段 / 不少于 8 张幻灯片 / ≥ 10 行」）→ 阈值键
_MIN_BLOCKS_RE = re.compile(
    r"(?:至少|不少于|≥|>=)\s*(\d+)\s*(?:个)?\s*(?P<unit>段落|段|行|张幻灯片|幻灯片|内容块)"
)
#: 内容块单位 → 非空阈值键（checklist 下限按单位细分，避免跨交付物误用）
_UNIT_KEYS: dict[str, str] = {
    "段落/表格": "段落",
    "幻灯片": "幻灯片",
    "数据行": "行",
    "文本行": "行",
}
#: 阈值键 → 正则里的单位词取值
_UNIT_ALIASES: dict[str, tuple[str, ...]] = {
    "段落": ("段落", "段", "内容块"),
    "幻灯片": ("幻灯片", "张幻灯片"),
    "行": ("行",),
}
#: 单位词 → 阈值键
_UNIT_BY_ALIAS: dict[str, str] = {
    alias: key for key, aliases in _UNIT_ALIASES.items() for alias in aliases
}
#: 引用 URL 数量（「引用至少 3 个来源 URL」）
_URL_COUNT_RE = re.compile(
    r"(?:至少|不少于|≥|>=)?\s*(\d+)\s*[个条处篇份]?\s*"
    r"(?:url|urls|链接|网址|来源|参考资料|参考文献|引用)",
    re.IGNORECASE,
)
#: 链接可达性要求（checklist 明确要求时才做 HEAD 探测）
_URL_STATUS_RE = re.compile(r"(?:head|可访问|能打开|有效性|状态码|200)", re.IGNORECASE)
#: 数据点（字面值）
_DATAPOINT_RE = re.compile(r"(?:数据点|数值|数据)\s*[:：]\s*(?P<value>.+)")
#: 显式正则判据
_REGEX_RE = re.compile(r"(?:正则|regex|模式)\s*[:：]\s*(?P<value>.+)")
#: 引号内关键词/章节名
_QUOTED_KEYWORD_RE = re.compile(r"[「『“\"']([^」』”\"'\n]{1,80})[」』”\"']")
#: 显式关键词清单
_KEYWORDS_RE = re.compile(r"(?:关键词|关键字|章节|小节|要点)\s*[:：]\s*(?P<value>.+)")
#: 命中数量下限
_MIN_HITS_RE = re.compile(r"(?:至少|不少于|命中)\s*(\d+)\s*个")
#: 命中率
_HIT_RATE_RE = re.compile(r"命中率\s*(?:≥|>=|至少)?\s*(\d+)\s*%")
#: 占位符相关词
_PLACEHOLDER_WORD_RE = re.compile(
    r"(占位符|占位|todo|tbd|undefined|待补充|待填写|空章节|模板变量|未填写)", re.IGNORECASE
)
#: 否定词（「不得包含…」）；否定类条目不做正向关键词抽取，避免把「不得含 X」判成「必须含 X」
_BAN_RE = re.compile(
    r"(不得|禁止|不能|不允许|避免|不含|不存在|没有|无\s*(?:占位|todo|tbd|undefined|待补充|空章节))",
    re.IGNORECASE,
)
#: URL 提取（排除中英文标点，避免多个 URL 被连成一个）
_URL_RE = re.compile(r"https?://[^\s<>\"'）)】\]，。；、（）]+")
#: 证据位置标记（L2 判分必须给出位置）
_EVIDENCE_LOCATOR_RE = re.compile(
    r"第\s*\d+|slide\s*\d+|paragraph\s*\d+|\d+\s*(?:段|页|张|行|句|节|章|部分|幻灯片|表格|图)",
    re.IGNORECASE,
)
#: 证据位置备选：引号内的标题名（如 标题「市场概览」）
_EVIDENCE_QUOTE_RE = re.compile(r"[「『“\"'][^」』”\"'\n]{2,60}[」』”\"']")


def _field(source: Any | None, name: str) -> str | None:
    """从 pydantic 模型或 dict 中取字符串字段（兼容两种入参形态）。"""
    if source is None:
        return None
    value = source.get(name) if isinstance(source, dict) else getattr(source, name, None)
    return value if isinstance(value, str) else None


def _clip(text: str, limit: int = _MAX_FIELD_CHARS) -> str:
    """截断上报字段（保留前缀，不静默丢信息）。"""
    text = (text or "").strip()
    return text if len(text) <= limit else text[:limit] + "…"


def _squash(text: str) -> str:
    """去除全部空白（判据与交付物文本比对时忽略换行/空格差异）。"""
    return re.sub(r"\s+", "", text or "")


def parse_checklist_items(checklist: str | None) -> list[str]:
    """解析 checklist 原文为条目列表（支持 ``-``/``*``/``1.``/``- [ ]``/表格/纯行）。"""
    lines = [line.strip() for line in (checklist or "").splitlines()]
    items: list[str] = []
    bullets = [_MARKUP_RE.sub("", m.group("text")) for m in map(_ITEM_LINE_RE.match, lines) if m]
    if bullets:
        candidates = bullets
    else:
        candidates = _table_or_plain_items(lines)
    for candidate in candidates:
        text = _MARKUP_RE.sub("", candidate).strip(" \t：:;；,，。")
        if text and text not in items:
            items.append(text)
    return items


def _table_or_plain_items(lines: list[str]) -> list[str]:
    """无项目符号时的兜底：Markdown 表格取最长单元格，否则取非空行。"""
    rows = [_TABLE_ROW_RE.match(line) for line in lines]
    if sum(1 for row in rows if row) >= 2:
        candidates: list[str] = []
        seen_body = False
        for row in rows:
            if row is None:
                continue
            cells = [cell.strip() for cell in row.group("cells").split("|")]
            cells = [cell for cell in cells if cell and not set(cell) <= set("-: ")]
            if not cells:
                continue
            candidate = max(cells, key=len)
            if not seen_body and _looks_like_table_header(candidate):
                seen_body = True  # 表头行不算验收条目，仅用于跳过
                continue
            seen_body = True
            candidates.append(candidate)
        return candidates
    return [line for line in lines if line and not line.startswith("#")]


def _looks_like_table_header(cell: str) -> bool:
    """表格首行是否为表头（「验收项/序号/说明」等，非验收要求本身）。"""
    return any(
        word in cell for word in ("验收项", "检查项", "序号", "编号", "说明", "备注", "状态")
    ) and "。" not in cell


def _parse_content_rules(
    items: list[ChecklistItem],
) -> tuple[list[ContentRule], list[ChecklistItem]]:
    """拆分 checklist 条目：可自动判定（L1）与需语义判分（L2）。

    只写结构下限（「数据表至少 3 行」）的条目由 L0 非空阈值消费，不再进 L1/L2，
    避免同一要求被重复判两次。
    """
    rules: list[ContentRule] = []
    semantic: list[ChecklistItem] = []
    for item in items:
        rule = _rule_of(item)
        if rule is not None:
            rules.append(rule)
        elif _MIN_BLOCKS_RE.search(item.text):
            continue
        else:
            semantic.append(item)
    return rules, semantic


def _rule_of(item: ChecklistItem) -> ContentRule | None:
    """单条 checklist → L1 判据（无法自动判定返回 None，交给 L2）。"""
    if _BAN_RE.search(item.text):
        # 否定类条目：只认「禁占位符」这一引擎级判据，其余交给 L2 语义判分
        if _PLACEHOLDER_WORD_RE.search(item.text):
            return ContentRule("placeholder", item.id, item.text)
        return None
    urls = _match_url_rule(item)
    if urls is not None:
        return urls
    datapoint = _match_datapoint_rule(item)
    if datapoint is not None:
        return datapoint
    return _match_keyword_rule(item)


def _match_url_rule(item: ChecklistItem) -> ContentRule | None:
    """「引用至少 N 个来源 URL」类判据。"""
    match = _URL_COUNT_RE.search(item.text)
    if match is None:
        return None
    count = int(match.group(1))
    if count <= 0 or count > 200:
        return None
    return ContentRule(
        "urls",
        item.id,
        item.text,
        min_count=count,
        check_url_status=bool(_URL_STATUS_RE.search(item.text)),
    )


def _match_datapoint_rule(item: ChecklistItem) -> ContentRule | None:
    """「数据点：X」/「正则：X」类判据（过长或正则非法则交给 L2）。"""
    match = _REGEX_RE.search(item.text)
    if match is not None:
        pattern = match.group("value").strip().strip("`").strip()
        try:
            re.compile(pattern)
        except re.error:
            return None
        return ContentRule(
            "datapoint", item.id, item.text, pattern=pattern, expected=pattern, literal=False
        )
    match = _DATAPOINT_RE.search(item.text)
    if match is None:
        return None
    value = match.group("value").strip().strip("`").strip("。；;，, ")
    if not value or len(value) > 60 or not any(char.isdigit() for char in value):
        return None
    return ContentRule(
        "datapoint",
        item.id,
        item.text,
        pattern=re.escape(_squash(value)),
        expected=value,
        literal=True,
    )


def _match_keyword_rule(item: ChecklistItem) -> ContentRule | None:
    """「包含「市场概览」章节」/「关键词：A、B」类判据。"""
    keywords = [text.strip() for text in _QUOTED_KEYWORD_RE.findall(item.text) if text.strip()]
    if not keywords:
        match = _KEYWORDS_RE.search(item.text)
        if match is None:
            return None
        keywords = [
            part.strip()
            for part in re.split(r"[、,，;；/]", match.group("value"))
            if part.strip()
        ]
    unique: list[str] = []
    for keyword in keywords:
        # 引号内的文件名属于交付物点名，不作为「正文必须出现」的关键词判据
        if keyword in unique or len(keyword) > 80 or _looks_like_file_name(keyword):
            continue
        unique.append(keyword)
    if not unique:
        return None
    return ContentRule(
        "keywords",
        item.id,
        item.text,
        keywords=unique,
        min_hits=_keyword_min_hits(item.text, len(unique)),
    )


def _keyword_min_hits(text: str, total: int) -> int:
    """关键词命中下限：默认全部命中，可被「至少 N 个」/「命中率 N%」放宽。"""
    match = _MIN_HITS_RE.search(text)
    if match is not None:
        return max(1, min(total, int(match.group(1))))
    rate = _HIT_RATE_RE.search(text)
    if rate is not None:
        percent = max(1, min(100, int(rate.group(1))))
        return max(1, math.ceil(total * percent / 100))
    return total


def _parse_deliverables(text: str) -> list[str]:
    """提取中枢点名的交付物文件名（只认引号内文件名，避免误判散文）。"""
    names: list[str] = []
    for match in _QUOTED_NAME_RE.finditer(text or ""):
        name = match.group("name").strip().lstrip("/")
        if name and name not in names:
            names.append(name)
    return names


def _looks_like_file_name(text: str) -> bool:
    """引号内文本是否为交付物文件名（不作为「必须出现在正文」的关键词）。"""
    return bool(
        re.search(
            r"\.(?:docx|docm|doc|pptx|ppt|xlsx|xlsm|xls|md|txt|csv|tsv|json|png|jpe?g|pdf)$",
            text.strip(),
            re.IGNORECASE,
        )
    )


def _parse_min_blocks(text: str) -> dict[str, int]:
    """解析按内容块单位细分的非空下限（缺省由引擎下限兜底）。"""
    thresholds: dict[str, int] = {}
    for value, unit in _MIN_BLOCKS_RE.findall(text or ""):
        key = _UNIT_BY_ALIAS.get(unit)
        count = int(value)
        if key is None or not 0 < count <= 200:
            continue
        thresholds[key] = max(thresholds.get(key, 0), count)
    return thresholds


# --------------------------------------------------------------------------- #
# 交付物读取（L0 结构闸的输入）
# --------------------------------------------------------------------------- #


@dataclass
class _Block:
    """交付物内容块（locator 即「证据位置」标注，如「第 3 段」）。"""

    locator: str
    text: str
    heading_level: int | None = None


@dataclass
class _Delivery:
    """一份交付物的解析结果（blocks 为空 + error 非空表示打不开）。"""

    name: str
    suffix: str
    reader: str = ""
    unit: str = ""
    blocks: list[_Block] = field(default_factory=list)
    error: str = ""
    #: 未知扩展名：只做存在性 + 字节非空校验，不做结构解析
    opaque: bool = False
    size: int = 0


def _docx_heading_level(paragraph: Any) -> int | None:
    """段落层级：Heading N → N；文档 Title/Subtitle → 0（标题不算章节，见下）。

    与 read_docx 的展示口径略有不同：这里把文档 Title（write_docx 的标题）
    视为 0 级，不参与「空章节」判定——标题后紧跟章节标题并非空章节缺陷。
    """
    name = str(getattr(getattr(paragraph, "style", None), "name", "") or "").strip()
    match = re.match(r"^Heading\s+(\d+)$", name)
    if match:
        return max(1, min(6, int(match.group(1))))
    if name.lower() in ("title", "subtitle"):
        return 0
    return None


def _read_docx(path: Path) -> tuple[str, str, list[_Block]]:
    """真打开 .docx（python-docx）并抽取段落/表格内容块。"""
    from docx import Document

    document = Document(str(path))
    blocks: list[_Block] = []
    index = 0
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        index += 1
        blocks.append(_Block(f"第 {index} 段", text, _docx_heading_level(paragraph)))
    for table_index, table in enumerate(document.tables, 1):
        cells = [cell.text.strip() for row in table.rows for cell in row.cells]
        text = " ".join(cell for cell in cells if cell)
        if text:
            blocks.append(_Block(f"第 {table_index} 张表", text))
    return "python-docx", "段落/表格", blocks


def _read_pptx(path: Path) -> tuple[str, str, list[_Block]]:
    """真打开 .pptx（python-pptx）并按幻灯片抽取内容块。

    幻灯片标题层级不参与「空章节」判定（PPT 的结构单位就是幻灯片，一页一标题
    并非缺陷），因此这里不标注 heading_level。
    """
    from pptx import Presentation

    presentation = Presentation(str(path))
    blocks: list[_Block] = []
    for index, slide in enumerate(presentation.slides, 1):
        parts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = (paragraph.text or "").strip()
                if text:
                    parts.append(text)
        if parts:
            blocks.append(_Block(f"第 {index} 张幻灯片", "\n".join(parts)))
    return "python-pptx", "幻灯片", blocks


def _read_xlsx(path: Path) -> tuple[str, str, list[_Block]]:
    """真打开 .xlsx（openpyxl）并按数据行抽取内容块。"""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        blocks: list[_Block] = []
        for sheet in workbook.worksheets:
            row_index = 0
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value).strip() for value in row if value not in (None, "")]
                if not cells:
                    continue
                row_index += 1
                blocks.append(_Block(f"「{sheet.title}」第 {row_index} 行", " | ".join(cells)))
        return "openpyxl", "数据行", blocks
    finally:
        workbook.close()


def _read_text(path: Path) -> tuple[str, str, list[_Block]]:
    """读取纯文本交付物（utf-8 → gbk 兜底），按非空行抽取内容块。"""
    data = path.read_bytes()
    text = ""
    for encoding in ("utf-8-sig", "gbk"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    blocks: list[_Block] = []
    index = 0
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        index += 1
        level = len(stripped) - len(stripped.lstrip("#"))
        blocks.append(_Block(f"第 {index} 行", stripped, level or None))
    return "文本读取", "文本行", blocks


#: 扩展名 → 真打开函数（L0「能被对应库真打开」的实现）
_READER_FUNCS: dict[str, Callable[[Path], tuple[str, str, list[_Block]]]] = {
    ".docx": _read_docx,
    ".pptx": _read_pptx,
    ".xlsx": _read_xlsx,
}


def _load_delivery(path: Path, work_dir: Path) -> _Delivery:
    """按扩展名真打开交付物；失败时用 error 记录条目级原因（不抛异常）。"""
    suffix = path.suffix.lower()
    name = _relative_name(path, work_dir)
    try:
        size = path.stat().st_size
    except OSError as exc:
        return _Delivery(name, suffix, error=f"无法读取文件信息：{type(exc).__name__}: {exc}")
    if suffix in OFFICE_READERS:
        reader, unit = OFFICE_READERS[suffix]
        try:
            actual_reader, actual_unit, blocks = _READER_FUNCS[suffix](path)
        except Exception as exc:  # noqa: BLE001 — 打不开按条目级失败记录（不终止任务）
            return _Delivery(
                name,
                suffix,
                reader=reader,
                unit=unit,
                size=size,
                error=f"无法打开（结构损坏或非有效{reader} 文档）：{type(exc).__name__}: {exc}",
            )
        return _Delivery(name, suffix, actual_reader, actual_unit, blocks, size=size)
    if suffix in TEXT_SUFFIXES:
        try:
            actual_reader, actual_unit, blocks = _read_text(path)
        except OSError as exc:
            return _Delivery(name, suffix, error=f"无法打开：{type(exc).__name__}: {exc}")
        return _Delivery(name, suffix, actual_reader, actual_unit, blocks, size=size)
    # 未知扩展名（png/pdf 等）：不做结构解析，只做存在性与字节非空校验
    return _Delivery(
        name,
        suffix,
        reader="字节存在性检查",
        unit="字节",
        opaque=True,
        size=size,
    )


def _relative_name(path: Path, work_dir: Path) -> str:
    """工作目录内相对路径（POSIX 风格）；越界时退回文件名。"""
    try:
        return path.resolve().relative_to(work_dir.resolve()).as_posix()
    except (OSError, ValueError):
        return path.name


# --------------------------------------------------------------------------- #
# L0 结构闸
# --------------------------------------------------------------------------- #


def scan_artifacts(work_dir: Path, *, since: float | None = None) -> list[str]:
    """扫描工作目录内可校验的交付物（相对 POSIX 路径，按名升序）。

    Args:
        work_dir: 任务工作目录。
        since: 只统计该时刻（epoch 秒）之后写入的文件，避免把同目录历史任务
            的旧产物当成本轮交付物；None 表示不限。
    """
    files, _truncated = scan_file_tree(work_dir)
    selected: list[str] = []
    for item in files:
        relative = str(item.get("path") or "")
        if not relative or is_excluded_relative_path(relative):
            continue
        if Path(relative).suffix.lower() not in ARTIFACT_SUFFIXES:
            continue
        if since is not None and float(item.get("mtime") or 0.0) < since - _MTIME_TOLERANCE:
            continue
        selected.append(relative)
    return selected


def _resolve_deliverable(work_dir: Path, name: str) -> tuple[Path | None, str]:
    """解析中枢点名的交付物路径（仅允许工作目录内相对路径）。"""
    try:
        return resolve_read_path(work_dir, name), ""
    except ToolError as exc:
        return None, str(exc).replace("素材", "交付物")


def _check_structure(
    work_dir: Path,
    plan: VerifyPlan,
    deliveries: dict[str, _Delivery],
    *,
    since: float | None = None,
) -> list[VerifyItem]:
    """L0 结构闸：交付物存在 → 真打开 → 非空（逐条产出条目级原因）。"""
    names = list(plan.deliverables) or scan_artifacts(work_dir, since=since)
    if not names:
        if plan.has_criteria:
            return [
                VerifyItem(
                    "L0-DELIVERY",
                    LAYER_STRUCTURE,
                    "交付物存在（中枢规格/checklist 要求产出交付物）",
                    False,
                    reason="工作目录内未发现任何交付物（docx/pptx/xlsx/md 等）",
                )
            ]
        return []
    items: list[VerifyItem] = []
    for index, name in enumerate(names, 1):
        prefix = f"L0-D{index}"
        path, error = _resolve_deliverable(work_dir, name)
        if path is None:
            items.append(
                VerifyItem(f"{prefix}-EXIST", LAYER_STRUCTURE, f"交付物存在：{name}", False, reason=error)
            )
            continue
        items.append(
            VerifyItem(
                f"{prefix}-EXIST", LAYER_STRUCTURE, f"交付物存在：{name}", True, evidence=name
            )
        )
        delivery = _load_delivery(path, work_dir)
        deliveries[delivery.name] = delivery
        items.append(_structure_open_item(prefix, name, delivery))
        items.append(
            _structure_nonempty_item(prefix, name, delivery, plan.min_blocks_for(delivery.unit))
        )
    return items


def _structure_open_item(prefix: str, name: str, delivery: _Delivery) -> VerifyItem:
    """「能被对应库真打开」条目。"""
    if delivery.opaque:
        return VerifyItem(
            f"{prefix}-OPEN",
            LAYER_STRUCTURE,
            f"交付物可被对应库打开：{name}",
            True,
            evidence=f"{delivery.reader}（该格式不做结构解析）",
        )
    if delivery.error:
        return VerifyItem(
            f"{prefix}-OPEN",
            LAYER_STRUCTURE,
            f"交付物可被 {delivery.reader} 真打开：{name}",
            False,
            reason=delivery.error,
        )
    return VerifyItem(
        f"{prefix}-OPEN",
        LAYER_STRUCTURE,
        f"交付物可被 {delivery.reader} 真打开：{name}",
        True,
        evidence=f"{delivery.reader} 打开成功（{len(delivery.blocks)} 个{delivery.unit}）",
    )


def _structure_nonempty_item(
    prefix: str, name: str, delivery: _Delivery, min_blocks: int
) -> VerifyItem:
    """「非空阈值」条目（未知格式按字节非空校验）。"""
    requirement = f"交付物非空（{delivery.unit}数 ≥ {min_blocks}）：{name}"
    if delivery.opaque:
        passed = delivery.size > 0
        return VerifyItem(
            f"{prefix}-NONEMPTY",
            LAYER_STRUCTURE,
            f"交付物非空（字节数 > 0）：{name}",
            passed,
            evidence=f"{delivery.size} 字节" if passed else "",
            reason="" if passed else "文件为 0 字节（空文件）",
        )
    if delivery.error:
        return VerifyItem(
            f"{prefix}-NONEMPTY",
            LAYER_STRUCTURE,
            requirement,
            False,
            reason=f"文件无法解析，无法统计{delivery.unit}：{delivery.error}",
        )
    count = len(delivery.blocks)
    passed = count >= min_blocks
    return VerifyItem(
        f"{prefix}-NONEMPTY",
        LAYER_STRUCTURE,
        requirement,
        passed,
        evidence=f"共 {count} 个{delivery.unit}（下限 {min_blocks}）" if passed else "",
        reason="" if passed else f"实际 {count} 个{delivery.unit}，低于下限 {min_blocks}（空文件或内容缺失）",
    )


# --------------------------------------------------------------------------- #
# L1 内容闸
# --------------------------------------------------------------------------- #


def _check_content(plan: VerifyPlan, deliveries: dict[str, _Delivery]) -> list[VerifyItem]:
    """L1 内容闸：禁占位符 / 禁空章节（引擎级）+ checklist 可判定项（任务级）。"""
    located = [
        (name, block) for name, delivery in deliveries.items() for block in delivery.blocks
    ]
    items = _check_placeholders(located, plan)
    items.extend(_check_empty_sections(deliveries))
    index = 0
    for rule in plan.content_rules:
        if rule.kind == "placeholder":
            continue  # 已由引擎级禁占位符条目覆盖（requirement 用 checklist 原文）
        index += 1
        items.append(_apply_rule(rule, index, located))
    return items


def _check_placeholders(
    located: list[tuple[str, _Block]], plan: VerifyPlan
) -> list[VerifyItem]:
    """禁占位符（TODO / undefined / {{…}} / 待补充 / 空章节占位）。"""
    explicit = next(
        (rule for rule in plan.content_rules if rule.kind == "placeholder"), None
    )
    item_id = f"L1-{explicit.item_id}" if explicit is not None else "L1-PLACEHOLDER"
    requirement = (
        explicit.requirement
        if explicit is not None
        else "交付物不得包含占位符（TODO / TBD / undefined / {{…}} / 待补充）"
    )
    hits: list[str] = []
    for label, pattern in PLACEHOLDER_PATTERNS:
        for name, block in located:
            match = re.search(pattern, block.text)
            if match is not None:
                hits.append(f"{label}「{match.group(0)[:40]}」（{name} {block.locator}）")
                break
    if not hits:
        return [
            VerifyItem(
                item_id,
                LAYER_CONTENT,
                requirement,
                True,
                evidence=f"已扫描 {len(located)} 个内容块，未发现占位符",
            )
        ]
    return [
        VerifyItem(
            item_id,
            LAYER_CONTENT,
            requirement,
            False,
            reason="命中占位符：" + "；".join(hits[:3]),
        )
    ]


def _check_empty_sections(deliveries: dict[str, _Delivery]) -> list[VerifyItem]:
    """禁空章节：标题（1 级及以上）下必须有正文内容。"""
    items: list[VerifyItem] = []
    index = 0
    for name, delivery in deliveries.items():
        headings = [block for block in delivery.blocks if _is_section_heading(block)]
        if not headings:
            continue
        index += 1
        empty = _empty_sections(delivery)
        if empty:
            items.append(
                VerifyItem(
                    f"L1-EMPTY{index}",
                    LAYER_CONTENT,
                    "不得存在空章节（标题下必须有正文内容）",
                    False,
                    reason=f"{name} 的空章节：" + "；".join(empty[:5]),
                )
            )
        else:
            items.append(
                VerifyItem(
                    f"L1-EMPTY{index}",
                    LAYER_CONTENT,
                    "不得存在空章节（标题下必须有正文内容）",
                    True,
                    evidence=f"{name}：{len(headings)} 个标题下均有正文内容",
                )
            )
    return items


def _is_section_heading(block: _Block) -> bool:
    """是否为章节标题（1 级及以上；文档 Title 记为 0 级不算章节）。"""
    return bool(block.heading_level) and block.heading_level >= 1


def _empty_sections(delivery: _Delivery) -> list[str]:
    """列出「标题后紧跟章节标题或直接结束」的空章节描述。"""
    descriptions: list[str] = []
    blocks = delivery.blocks
    for position, block in enumerate(blocks):
        if not _is_section_heading(block):
            continue
        following = blocks[position + 1] if position + 1 < len(blocks) else None
        if following is None or _is_section_heading(following):
            descriptions.append(f"「{block.text[:40]}」（{block.locator}）")
    return descriptions


def _apply_rule(rule: ContentRule, index: int, located: list[tuple[str, _Block]]) -> VerifyItem:
    """执行单条 L1 判据并给出证据位置或未通过原因。"""
    item_id = f"L1-{rule.item_id}" if rule.item_id else f"L1-R{index}"
    if rule.kind == "urls":
        return _check_urls(rule, item_id, located)
    if rule.kind == "datapoint":
        return _check_datapoint(rule, item_id, located)
    return _check_keywords(rule, item_id, located)


def _check_keywords(rule: ContentRule, item_id: str, located: list[tuple[str, _Block]]) -> VerifyItem:
    """必需章节/关键词命中率判定（忽略空白差异）。"""
    matched: list[str] = []
    missing: list[str] = []
    for keyword in rule.keywords:
        needle = _squash(keyword)
        hit = next(
            (
                f"「{keyword}」（{name} {block.locator}）"
                for name, block in located
                if needle and needle in _squash(block.text)
            ),
            "",
        )
        if hit:
            matched.append(hit)
        else:
            missing.append(f"「{keyword}」")
    passed = len(matched) >= rule.min_hits
    if passed:
        return VerifyItem(
            item_id,
            LAYER_CONTENT,
            rule.requirement,
            True,
            evidence=f"命中 {len(matched)}/{len(rule.keywords)}：{'、'.join(matched[:5])}",
        )
    detail = f"仅命中 {len(matched)}/{len(rule.keywords)} 个必需章节/关键词（下限 {rule.min_hits}）"
    if missing:
        detail += "；缺失：" + "、".join(missing[:8])
    if not located:
        detail += "；交付物无可用内容块"
    return VerifyItem(item_id, LAYER_CONTENT, rule.requirement, False, reason=detail)


def _check_datapoint(rule: ContentRule, item_id: str, located: list[tuple[str, _Block]]) -> VerifyItem:
    """必需数据点判定（字面值按去空白比对，显式正则按原文比对）。"""
    for name, block in located:
        haystack = _squash(block.text) if rule.literal else block.text
        match = re.search(rule.pattern, haystack)
        if match is not None:
            return VerifyItem(
                item_id,
                LAYER_CONTENT,
                rule.requirement,
                True,
                evidence=f"命中数据点「{match.group(0)[:40]}」（{name} {block.locator}）",
            )
    return VerifyItem(
        item_id,
        LAYER_CONTENT,
        rule.requirement,
        False,
        reason=f"未找到必需数据点：{rule.expected}",
    )


def _check_urls(rule: ContentRule, item_id: str, located: list[tuple[str, _Block]]) -> VerifyItem:
    """引用 URL 数量判定（可选可达性由 ``run_verify(probe_urls=True)`` 追加）。"""
    urls = _extract_urls(located)
    passed = len(urls) >= rule.min_count
    if passed:
        return VerifyItem(
            item_id,
            LAYER_CONTENT,
            rule.requirement,
            True,
            evidence=f"引用 {len(urls)} 个 URL（下限 {rule.min_count}）：{'、'.join(urls[:3])}",
        )
    return VerifyItem(
        item_id,
        LAYER_CONTENT,
        rule.requirement,
        False,
        reason=f"实际引用 {len(urls)} 个 URL，低于要求 ≥ {rule.min_count}",
    )


def _extract_urls(located: list[tuple[str, _Block]]) -> list[str]:
    """按出现顺序提取去重 URL。"""
    urls: list[str] = []
    for _name, block in located:
        for url in _URL_RE.findall(block.text):
            clean = url.rstrip(".,;，。；、")
            if clean and clean not in urls:
                urls.append(clean)
    return urls


async def _probe_urls(urls: list[str]) -> list[str]:
    """HEAD 探测引用链接可达性（失败回退 GET）；返回不可达 URL 描述。"""
    unreachable: list[str] = []
    try:
        async with httpx.AsyncClient(timeout=_URL_PROBE_TIMEOUT, follow_redirects=True) as client:
            for url in urls:
                try:
                    response = await client.head(url)
                    if response.status_code >= 400:
                        response = await client.get(url)
                    if response.status_code >= 400:
                        unreachable.append(f"{url}（HTTP {response.status_code}）")
                except httpx.HTTPError:
                    unreachable.append(f"{url}（无法访问）")
    except httpx.HTTPError as exc:  # 探测设施本身不可用：不据此判失败，仅记录
        logger.warning("verify_url_probe_unavailable", extra={"error": str(exc)})
    return unreachable


async def _check_url_status(
    plan: VerifyPlan, located: list[tuple[str, _Block]]
) -> list[VerifyItem]:
    """checklist 明确要求链接可访问时，追加 HEAD 200 可达性条目。"""
    items: list[VerifyItem] = []
    for rule in plan.content_rules:
        if rule.kind != "urls" or not rule.check_url_status:
            continue
        urls = _extract_urls(located)
        unreachable = await _probe_urls(urls)
        items.append(
            VerifyItem(
                f"L1-{rule.item_id}-HTTP",
                LAYER_CONTENT,
                rule.requirement,
                not unreachable,
                evidence=f"{len(urls)} 个引用链接均可访问（HEAD 200）" if not unreachable else "",
                reason="" if not unreachable else "不可访问的引用链接：" + "；".join(unreachable[:5]),
            )
        )
    return items


# --------------------------------------------------------------------------- #
# L2 语义闸（LLM 逐条判 + 必须给证据位置）
# --------------------------------------------------------------------------- #


def has_evidence_location(evidence: str) -> bool:
    """判断 L2 证据是否给出**位置**（第几段/第几张幻灯片/引号内标题）。

    ADR-065 M4：只给结论不给证据位置视为不通过。
    """
    text = (evidence or "").strip()
    if not text:
        return False
    return bool(_EVIDENCE_LOCATOR_RE.search(text) or _EVIDENCE_QUOTE_RE.search(text))


def _render_deliveries(deliveries: dict[str, _Delivery], limit: int = _L2_MAX_CONTENT_CHARS) -> str:
    """渲染交付物正文供判分（带证据位置标记，超长截断并标注）。"""
    lines: list[str] = []
    total = 0
    for name, delivery in deliveries.items():
        if delivery.error or delivery.opaque:
            lines.append(f"【交付物】{name}（无法读取内容：{delivery.error or '该格式不做解析'}）")
            continue
        header = f"【交付物】{name}（{len(delivery.blocks)} 个{delivery.unit}）"
        lines.append(header)
        total += len(header)
        for block in delivery.blocks:
            marker = "[标题] " if block.heading_level else ""
            line = f"[{block.locator}]{marker}{block.text}"
            if total + len(line) > limit:
                lines.append("…（内容过长已截断，请按已展示内容判分）")
                return "\n".join(lines)
            lines.append(line)
            total += len(line)
    return "\n".join(lines) or "（无交付物内容）"


def _build_semantic_prompt(items: list[ChecklistItem], deliveries: dict[str, _Delivery]) -> str:
    """构造 L2 判分提示词（逐条判 + 必须给证据位置 + 仅输出 JSON）。"""
    checklist = "\n".join(f"{item.id}. {item.text}" for item in items)
    return (
        "你是办公交付物验收判分员，请按验收清单逐条判断交付物是否达标。\n\n"
        f"【验收清单】\n{checklist}\n\n"
        f"【交付物内容（行首方括号内是证据位置标记）】\n{_render_deliveries(deliveries)}\n\n"
        "判分要求：\n"
        "1. 只依据上面展示的交付物内容判断，不要臆测未展示的内容；\n"
        "2. 每条必须给出 evidence 字段写明证据位置（如「第 3 段」「第 2 张幻灯片」"
        "「标题「市场概览」」）；**只给结论不给证据位置的条目一律按不通过处理**；\n"
        "3. 仅输出 JSON，不要输出任何解释文字，格式：\n"
        '{"items":[{"id":"C1","passed":true,"evidence":"第 3 段：……","reason":""}]}'
    )


def _parse_semantic_response(raw: str) -> dict[str, dict[str, Any]]:
    """解析判分输出（容忍 ```json 围栏与前后解释文字）；返回 id → 条目。"""
    text = (raw or "").strip()
    fenced = re.search(r"```(?:json)?\s*(?P<body>.+?)```", text, re.DOTALL)
    if fenced is not None:
        text = fenced.group("body").strip()
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        return {}
    try:
        data = json.loads(text[start : end + 1])
    except ValueError:
        return {}
    entries = data.get("items") if isinstance(data, dict) else None
    if not isinstance(entries, list):
        return {}
    judged: dict[str, dict[str, Any]] = {}
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        key = str(entry.get("id") or "").strip()
        if key:
            judged[key] = entry
    return judged


def _semantic_failures(items: list[ChecklistItem], reason: str) -> list[VerifyItem]:
    """L2 无法执行时统一按未通过计（不许假成功）。"""
    return [
        VerifyItem(f"L2-{item.id}", LAYER_SEMANTIC, item.text, False, reason=reason)
        for item in items
    ]


async def _check_semantic(
    plan: VerifyPlan, deliveries: dict[str, _Delivery], llm: Any | None
) -> list[VerifyItem]:
    """L2 语义闸：LLM 逐条判分，证据位置缺失/缺失结论/调用失败一律未通过。"""
    if not plan.semantic:
        return []
    if llm is None:
        return _semantic_failures(plan.semantic, "未提供判分模型，L2 语义闸无法执行")
    messages = [{"role": "user", "content": _build_semantic_prompt(plan.semantic, deliveries)}]
    try:
        response = await llm.chat(messages, None)
    except LLMError as exc:
        return _semantic_failures(plan.semantic, f"L2 语义判分不可用（LLM 调用失败）：{exc}")
    except Exception as exc:  # noqa: BLE001 — 判分客户端异常不得中断任务，按未通过计
        return _semantic_failures(
            plan.semantic, f"L2 语义判分不可用（{type(exc).__name__}: {exc}）"
        )
    message = response.get("message") if isinstance(response, dict) else None
    raw = str((message or {}).get("content") or "")
    judged = _parse_semantic_response(raw)
    items: list[VerifyItem] = []
    for entry in plan.semantic:
        items.append(_semantic_item(entry, judged.get(entry.id)))
    return items


def _semantic_item(entry: ChecklistItem, judged: dict[str, Any] | None) -> VerifyItem:
    """单条 L2 结论 → 条目（证据位置缺失视为不通过）。"""
    item_id = f"L2-{entry.id}"
    if judged is None:
        return VerifyItem(item_id, LAYER_SEMANTIC, entry.text, False, reason="判分输出缺少该条结论")
    passed = bool(judged.get("passed"))
    evidence = str(judged.get("evidence") or "").strip()
    reason = str(judged.get("reason") or "").strip()
    if passed and not has_evidence_location(evidence):
        detail = (
            "只给结论未给证据位置（要求注明第几段/第几张幻灯片）"
            if not evidence
            else f"证据未给出可核验的位置（收到「{_clip(evidence, 60)}」；要求如「第 3 段」）"
        )
        return VerifyItem(item_id, LAYER_SEMANTIC, entry.text, False, evidence=evidence, reason=detail)
    if passed:
        return VerifyItem(item_id, LAYER_SEMANTIC, entry.text, True, evidence=evidence)
    return VerifyItem(
        item_id, LAYER_SEMANTIC, entry.text, False, reason=reason or "判分模型判定未满足该条"
    )


# --------------------------------------------------------------------------- #
# 入口：单次判定与回灌闸门
# --------------------------------------------------------------------------- #


async def run_verify(
    work_dir: Path,
    plan: VerifyPlan,
    *,
    llm: Any | None = None,
    since: float | None = None,
    probe_urls: bool = False,
) -> VerifyReport:
    """执行三层判据（L0 → L1 → L2，任一层不过即短路）。

    Args:
        work_dir: 任务工作目录（交付物所在目录）。
        plan: 判据计划（只来自中枢下发的规格/checklist）。
        llm: 判分用 LLM 客户端（需提供 ``chat(messages, tools=None)``）。
        since: 只把该时刻之后写入的文件视为本轮交付物（避免历史产物干扰）。
        probe_urls: 是否对 checklist 要求「可访问」的引用链接做 HEAD 探测。

    Returns:
        ``VerifyReport``：``passed`` + 首个未通过层 ``layer`` + 逐条 ``items``。
    """
    deliveries: dict[str, _Delivery] = {}
    structure_items = await asyncio.to_thread(
        _check_structure, work_dir, plan, deliveries, since=since
    )
    if not structure_items and not plan.has_criteria:
        return VerifyReport(
            passed=True,
            layer=LAYER_PASSED,
            source=plan.source,
            note="中枢未下发判据且未发现交付物，交付质量闸未生效",
        )
    items = structure_items
    if _has_failure(structure_items):
        return VerifyReport(False, LAYER_STRUCTURE, items, plan.source)
    located = [
        (name, block) for name, delivery in deliveries.items() for block in delivery.blocks
    ]
    content_items: list[VerifyItem] = []
    if probe_urls:
        content_items.extend(await _check_url_status(plan, located))
    content_items.extend(await asyncio.to_thread(_check_content, plan, deliveries))
    items = items + content_items
    if _has_failure(content_items):
        return VerifyReport(False, LAYER_CONTENT, items, plan.source)
    semantic_items = await _check_semantic(plan, deliveries, llm)
    items = items + semantic_items
    if _has_failure(semantic_items):
        return VerifyReport(False, LAYER_SEMANTIC, items, plan.source)
    return VerifyReport(True, LAYER_PASSED, items, plan.source)


def _has_failure(items: list[VerifyItem]) -> bool:
    """层内是否存在未通过条目。"""
    return any(not item.passed for item in items)


async def run_verify_loop(
    work_dir: Path,
    plan: VerifyPlan,
    *,
    llm: Any | None,
    repair: Callable[[str, int], Awaitable[Any]],
    on_round: Callable[[VerifyReport, str], Awaitable[None]] | None = None,
    since: float | None = None,
    probe_urls: bool = False,
    max_repair_rounds: int = MAX_REPAIR_ROUNDS,
) -> VerifyReport:
    """产出后跑三层闸；不过 → **只回灌未通过条目**重做，≤5 轮仍不过即返回（调用方阻断）。

    Args:
        work_dir: 任务工作目录。
        plan: 判据计划。
        llm: 判分用 LLM 客户端。
        repair: 回灌回调 ``(instruction, round_index) -> Awaitable``；instruction
            只含未通过条目 + 期望 + 实际，已达标项不出现。
        on_round: 每轮判定回调 ``(report, phase)``，phase 为 ``initial`` /
            ``repair_N``（进度上报用）。
        since: 交付物时间下界（任务起始时刻）。
        probe_urls: 是否探测引用链接可达性。
        max_repair_rounds: 回灌轮上限（默认 5，达到后不再发起下一轮）。
    """
    report = await run_verify(work_dir, plan, llm=llm, since=since, probe_urls=probe_urls)
    if on_round is not None:
        await on_round(report, "initial")
    for round_index in range(1, max_repair_rounds + 1):
        if report.passed:
            return report
        instruction = report.repair_instruction(round_index, max_repair_rounds)
        await repair(instruction, round_index)
        report = await run_verify(work_dir, plan, llm=llm, since=since, probe_urls=probe_urls)
        report.repair_round = round_index
        if on_round is not None:
            await on_round(report, f"repair_{round_index}")
    return report
