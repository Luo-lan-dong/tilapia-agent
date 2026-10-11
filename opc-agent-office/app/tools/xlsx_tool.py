"""表格写出与另存工具（ADR-065 M2）：write_xlsx / edit_xlsx。

任务常需产出 Excel 表格，或在既有素材表上追加行、改单元格。本模块提供两个工具：

- ``write_xlsx``：按 ``sheets``（或 ``headers`` + ``rows``）新建多工作表工作簿；
- ``edit_xlsx``：读入既有工作簿 → 应用 ``ops`` → **另存为新文件**（原件不动）。

安全红线（与 ADR-065 M1 只读工具一致，读入侧直接复用 ``resolve_read_path``）：

- 输入文件仅允许工作目录内已存在的相对路径，绝对路径与 ``..`` 越界一律 ``ToolError``；
- 输出走基类 ``_resolve_output``（仅允许工作目录内相对路径）；
- ``edit_xlsx`` 的输出名与输入名解析后相同 → ``ToolError`` 拒绝（保护原素材，
  不做原地覆盖，避免与 D1 冲突）。
"""

import logging
import re
from copy import copy
from pathlib import Path
from typing import Any

from app.tools.base import OfficeTool, ToolError, clean_llm_text, require_str
from app.tools.read_tool import resolve_read_path

logger = logging.getLogger(__name__)

#: 可选工作表名缺省值（与 write_docx 的 title 不同，这里允许完全缺省）
DEFAULT_SHEET_NAME = "Sheet1"

#: 新建工作簿时的默认工作表数上限（防 LLM 传入超大数组）
_MAX_SHEETS = 50

#: 单元格写入的字符上限（超长截断，防单格塞入整篇正文）
_MAX_CELL_CHARS = 2_000

#: edit_xlsx 支持的操作类型
_SUPPORTED_OPS = ("append_rows", "set_cell", "add_sheet")

#: Excel 单元格地址（如 A1、AB12），用于 set_cell 入参校验
_CELL_RE = re.compile(r"^[A-Za-z]{1,3}[1-9][0-9]{0,6}$")

#: Excel 工作表名非法字符（openpyxl 会直接抛异常，这里提前转可读错误）
_INVALID_SHEET_CHARS = (":", "\\", "/", "?", "*", "[", "]")


def _require_xlsx_suffix(target: Path, tool_name: str) -> None:
    """校验扩展名与工具匹配（格式不符 → 可读错误，提示改用对应工具）。"""
    if target.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ToolError(
            f"{tool_name} 仅支持 .xlsx/.xlsm 文件，收到：{target.name}"
            "（请确认文件格式；只读素材请改用 read_xlsx）"
        )


def _safe_sheet_name(name: Any, *, default: str | None = DEFAULT_SHEET_NAME) -> str:
    """校验并规整工作表名（openpyxl 限制：非空、≤31 字符、不含 : \\ / ? * [ ]）。

    ``default=None`` 表示目标工作表必须显式给出（如 append_rows/set_cell）。
    """
    raw = clean_llm_text(str(name or "").strip()) or (default or "")
    if not raw:
        raise ToolError("必须指定工作表名 sheet")
    if len(raw) > 31:
        raise ToolError(f"工作表名超长（最长 31 字符）：{raw[:40]}")
    if any(char in raw for char in _INVALID_SHEET_CHARS):
        raise ToolError(f"工作表名含非法字符（: \\ / ? * [ ]）：{raw}")
    return raw


def _coerce_cell(value: Any) -> Any:
    """单元格值清洗：LLM 长文本截断；其余标量原样写入以保留类型。"""
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = clean_llm_text(str(value))
    return text[:_MAX_CELL_CHARS]


def _coerce_table_rows(rows: Any, *, param: str, allow_empty: bool = False) -> list[list[Any]]:
    """校验二维数组入参（行必须是数组，单元格可缺省）。"""
    if not isinstance(rows, list) or (not rows and not allow_empty):
        raise ToolError(f"参数 {param} 必须为非空二维数组")
    table: list[list[Any]] = []
    for index, row in enumerate(rows):
        if not isinstance(row, list):
            raise ToolError(f"参数 {param} 第 {index + 1} 行必须为数组")
        table.append([_coerce_cell(cell) for cell in row])
    return table


class WriteXlsxTool(OfficeTool):
    """生成 .xlsx 工作簿（多工作表 + 可选表头加粗）。"""

    name = "write_xlsx"
    description = (
        "生成 Excel 工作簿（.xlsx）。用 sheets 描述工作表：每个元素含 name（表名）与 rows"
        "（二维数组，首行视为表头）；也可用简写 headers + rows 只写单表。"
        "表头默认加粗（bold_header=false 可关闭）。文件名以 .xlsx 结尾，保存在工作目录内。"
    )
    parameters: dict[str, Any] = {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "文件名，如：数据汇总.xlsx"},
            "sheets": {
                "type": "array",
                "description": "工作表列表，按顺序创建",
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string", "description": "工作表名（缺省 Sheet1）"},
                        "rows": {
                            "type": "array",
                            "items": {"type": "array"},
                            "description": "二维单元格数组，首行为表头",
                        },
                    },
                    "required": ["rows"],
                },
            },
            "headers": {
                "type": "array",
                "description": "单表简写：表头列名（与 rows 搭配，等价于 sheets 里的第一行）",
                "items": {"type": "string"},
            },
            "rows": {
                "type": "array",
                "description": "单表简写：数据行二维数组（配合 headers 使用时不含表头）",
                "items": {"type": "array"},
            },
            "bold_header": {
                "type": "boolean",
                "description": "是否加粗首行表头（默认 true）",
            },
        },
        "required": ["filename"],
    }

    def execute(self, args: dict[str, Any]) -> str:
        """渲染 sheets 为 xlsx 并落盘。"""
        try:
            from openpyxl import Workbook  # 延迟导入：仅执行时需要
        except ImportError as exc:  # pragma: no cover - 环境缺依赖
            raise ToolError(f"openpyxl 未安装: {exc}") from exc

        filename = require_str(args, "filename", max_len=200)
        bold_header = args.get("bold_header")
        if bold_header is not None and not isinstance(bold_header, bool):
            raise ToolError("参数 bold_header 必须为布尔值")
        do_bold = True if bold_header is None else bold_header

        sheets = self._collect_sheets(args)
        raw_name = filename if filename.lower().endswith(".xlsx") else f"{filename}.xlsx"
        target = self._resolve_output(raw_name)
        _require_xlsx_suffix(target, self.name)

        workbook = Workbook()
        workbook.remove(workbook.active)
        written = 0
        for name, rows in sheets:
            sheet = workbook.create_sheet(title=name)
            for row in rows:
                sheet.append(row)
            if do_bold and rows:
                for cell in sheet[1]:
                    # openpyxl 样式不可变：必须复制后整体赋值，不能改 cell.font.bold
                    font = copy(cell.font)
                    font.bold = True
                    cell.font = font
            written += len(rows)
        workbook.save(str(target))
        logger.info(
            "tool_xlsx_written",
            extra={
                "event": "tool_xlsx_written",
                "file": str(target),
                "sheets": len(sheets),
                "rows": written,
            },
        )
        names = "、".join(name for name, _ in sheets)
        return (
            f"已生成 {target.name}（{len(sheets)} 个工作表：{names}，共 {written} 行，"
            f"{target.stat().st_size} 字节）"
        )

    def _collect_sheets(self, args: dict[str, Any]) -> list[tuple[str, list[list[Any]]]]:
        """归一 sheets / headers+rows 两种入参形式，校验后返回 [(表名, 数据行)]。"""
        raw_sheets = args.get("sheets")
        sheets: list[tuple[str, list[list[Any]]]] = []
        if isinstance(raw_sheets, list) and raw_sheets:
            if len(raw_sheets) > _MAX_SHEETS:
                raise ToolError(f"sheets 数量超限（最多 {_MAX_SHEETS} 个工作表）")
            for index, item in enumerate(raw_sheets):
                if not isinstance(item, dict):
                    raise ToolError(f"sheets 第 {index + 1} 项必须为对象")
                name = _safe_sheet_name(item.get("name") or DEFAULT_SHEET_NAME)
                if name in [existing for existing, _ in sheets]:
                    raise ToolError(f"工作表名重复：{name}")
                rows = _coerce_table_rows(item.get("rows"), param=f"sheets[{index + 1}].rows")
                sheets.append((name, rows))
            return sheets

        headers = args.get("headers")
        rows = args.get("rows")
        if isinstance(headers, list) and headers:
            table = [[_coerce_cell(cell) for cell in headers]]
            if rows is not None:
                table.extend(_coerce_table_rows(rows, param="rows"))
            return [(_safe_sheet_name(args.get("sheet") or DEFAULT_SHEET_NAME), table)]
        if isinstance(rows, list) and rows:
            name = _safe_sheet_name(args.get("sheet") or DEFAULT_SHEET_NAME)
            return [(name, _coerce_table_rows(rows, param="rows"))]
        raise ToolError("必须提供非空的 sheets，或 headers + rows（或单独的 rows）")


class EditXlsxTool(OfficeTool):
    """在既有 .xlsx 上做追加行 / 改单元格 / 加表，并另存为新文件。"""

    name = "edit_xlsx"
    description = (
        "编辑工作目录内已有的 Excel 工作簿（.xlsx）并**另存为新文件**（原文件保持不变）。"
        "ops 支持：append_rows（sheet + rows，向已有工作表末尾追加行）、"
        "set_cell（sheet + cell + value，改写指定单元格）、add_sheet（name + rows，新增工作表）。"
        "output_filename 缺省时自动派生命名（如 表.xlsx → 表-edited.xlsx），返回文本会说明实际输出名。"
    )
    parameters: dict[str, Any] = {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "已存在的相对路径，如：数据/销量表.xlsx"},
            "ops": {
                "type": "array",
                "description": "编辑操作列表（按顺序应用）",
                "items": {
                    "type": "object",
                    "properties": {
                        "op": {
                            "type": "string",
                            "enum": list(_SUPPORTED_OPS),
                            "description": "操作类型",
                        },
                        "sheet": {"type": "string", "description": "目标工作表名"},
                        "rows": {
                            "type": "array",
                            "items": {"type": "array"},
                            "description": "append_rows / add_sheet 的二维数据行",
                        },
                        "cell": {"type": "string", "description": "set_cell 的单元格地址，如 B2"},
                        "value": {"description": "set_cell 写入的值（文本或数值）"},
                        "name": {"type": "string", "description": "add_sheet 的新工作表名"},
                    },
                    "required": ["op"],
                },
            },
            "output_filename": {
                "type": "string",
                "description": "可选：另存文件名（必须与 filename 不同，缺省自动加 -edited）",
            },
        },
        "required": ["filename", "ops"],
    }

    def execute(self, args: dict[str, Any]) -> str:
        """读入 → 改 → 另存；失败抛 ``ToolError``（不抛未捕获异常）。"""
        try:
            from openpyxl import load_workbook  # 延迟导入：仅执行时需要
        except ImportError as exc:  # pragma: no cover - 环境缺依赖
            raise ToolError(f"openpyxl 未安装: {exc}") from exc

        filename = require_str(args, "filename", max_len=500)
        source = resolve_read_path(self.work_dir, filename)
        _require_xlsx_suffix(source, self.name)
        ops = args.get("ops")
        if not isinstance(ops, list) or not ops:
            raise ToolError("ops 必须为非空数组")
        target = self.resolve_edit_output(
            input_path=source,
            output_filename=args.get("output_filename"),
            suffix=".xlsx",
        )

        try:
            workbook = load_workbook(str(source))
        except Exception as exc:  # noqa: BLE001 — 解析异常转可读错误，不终止工具循环
            logger.warning(
                "tool_edit_load_failed",
                extra={"event": "tool_edit_load_failed", "tool": self.name, "error": str(exc)},
            )
            raise ToolError(
                f"{self.name} 读取失败（{source.name}）：{type(exc).__name__}: {exc}"
            ) from exc
        try:
            applied = [self._apply_op(workbook, op, index) for index, op in enumerate(ops)]
            workbook.save(str(target))
        finally:
            workbook.close()
        logger.info(
            "tool_xlsx_edited",
            extra={
                "event": "tool_xlsx_edited",
                "tool": self.name,
                "source": str(source),
                "output": str(target),
                "ops": len(applied),
            },
        )
        return (
            f"已另存为新文件 {target.name}（源文件 {source.name} 保持不变），"
            f"共应用 {len(applied)} 个操作：\n" + "\n".join(f"- {item}" for item in applied)
        )

    def _apply_op(self, workbook: Any, op: Any, index: int) -> str:
        """应用单个操作并返回可读描述；非法操作抛 ``ToolError``（不静默跳过）。"""
        if not isinstance(op, dict):
            raise ToolError(f"ops 第 {index + 1} 项必须为对象")
        kind = str(op.get("op") or "").strip()
        if kind not in _SUPPORTED_OPS:
            raise ToolError(
                f"ops 第 {index + 1} 项 op 不支持：{kind or '(空)'}（可选：{'/'.join(_SUPPORTED_OPS)}）"
            )
        if kind == "add_sheet":
            name = _safe_sheet_name(op.get("name") or op.get("sheet"))
            if name in workbook.sheetnames:
                raise ToolError(f"工作表已存在，无法新建：{name}（改用 append_rows 追加）")
            rows = _coerce_table_rows(op.get("rows"), param="rows")
            sheet = workbook.create_sheet(title=name)
            for row in rows:
                sheet.append(row)
            return f"新建工作表 {name}（{len(rows)} 行）"

        name = _safe_sheet_name(op.get("sheet"), default=None)
        if name not in workbook.sheetnames:
            raise ToolError(f"工作表不存在：{name}（现有：{'、'.join(workbook.sheetnames)}）")
        sheet = workbook[name]
        if kind == "append_rows":
            rows = _coerce_table_rows(op.get("rows"), param="rows")
            for row in rows:
                sheet.append(row)
            return f"向工作表 {name} 追加 {len(rows)} 行"
        cell = str(op.get("cell") or "").strip()
        if not _CELL_RE.match(cell):
            raise ToolError(f"单元格地址非法：{cell or '(空)'}（示例：B2、AA10）")
        if "value" not in op:
            raise ToolError("set_cell 缺少 value 参数")
        sheet[cell.upper()] = _coerce_cell(op.get("value"))
        return f"改写 {name}!{cell.upper()}"
