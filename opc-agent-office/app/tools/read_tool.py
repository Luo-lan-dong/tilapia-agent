"""素材读取工具（ADR-065 M1）：read_docx / read_xlsx / read_csv / read_text。

任务常附带素材文件（工作目录内的 docx/xlsx/csv/txt/md 等），LLM 产出前须先
读取素材真实内容。本模块提供四个**只读**工具，统一遵守安全红线：

- 仅允许工作目录内的相对路径；绝对路径与 ``..`` 越界一律 ``ToolError``；
- 文件不存在 / 扩展名与工具不符 → ``ToolError`` 带可读原因（回传 LLM 自行调整）；
- 单次返回文本上限 ``MAX_OUTPUT_CHARS``，超出截断并标注 ``truncated``。

与写工具（write_docx/write_pptx）不同，本模块只读不写：不创建目录、不改动素材。
"""

import csv
import io
import logging
import re
from abc import abstractmethod
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app.tools.base import OfficeTool, ToolError, require_str

logger = logging.getLogger(__name__)

#: 单次返回给 LLM 的文本上限（超出截断并标注，防撑爆上下文）
MAX_OUTPUT_CHARS = 20_000

#: read_csv 默认读取的数据行数（不含表头）
DEFAULT_CSV_ROWS = 50

#: read_xlsx 单个 sheet 默认读取的最大行数
DEFAULT_SHEET_ROWS = 500

#: 行数参数硬上限（防止 LLM 传入超大值）
_HARD_ROW_CAP = 5_000

#: read_text 允许的纯文本扩展名
_TEXT_EXTENSIONS = (
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".log",
    ".yaml",
    ".yml",
    ".ini",
    ".cfg",
    ".conf",
    ".csv",
    ".tsv",
    ".rst",
    ".xml",
    ".html",
    ".htm",
    ".sql",
    ".py",
    ".sh",
    ".toml",
)

#: Word 标题样式名（英文模板 "Heading 1"，中文模板同样映射该样式名）
_HEADING_RE = re.compile(r"^Heading\s+(\d+)$")


def truncate_output(text: str, *, limit: int = MAX_OUTPUT_CHARS) -> str:
    """截断过长文本并标注 ``truncated``（保留可读提示，不静默丢内容）。"""
    if len(text) <= limit:
        return text
    return (
        text[:limit]
        + f"\n\n[[truncated]] 内容超过 {limit} 字符上限，已截断；"
        "如需后续部分请缩小范围（如指定 sheet、减少行数，或分段读取）。"
    )


def resolve_read_path(work_dir: Path, filename: str) -> Path:
    """解析只读素材路径：仅允许工作目录内已存在的相对路径（防越界读取）。"""
    if not filename or not filename.strip():
        raise ToolError("filename 不能为空")
    raw = filename.strip().replace("\\", "/")
    if raw.startswith("/") or re.match(r"^[A-Za-z]:", raw):
        raise ToolError(f"仅允许工作目录内的相对路径，收到绝对路径：{filename}")
    parts = [part for part in raw.split("/") if part not in ("", ".")]
    if any(part == ".." for part in parts):
        raise ToolError(f"路径越出工作目录（不允许 ..）：{filename}")
    if not parts:
        raise ToolError(f"非法文件名：{filename}")
    root = work_dir.resolve()
    resolved = (work_dir.joinpath(*parts)).resolve()
    if not resolved.is_relative_to(root):
        raise ToolError(f"路径越出工作目录：{filename}")
    if not resolved.exists():
        raise ToolError(f"素材文件不存在：{filename}")
    if not resolved.is_file():
        raise ToolError(f"素材路径不是文件：{filename}")
    return resolved


def _require_suffix(target: Path, allowed: tuple[str, ...], tool_name: str) -> None:
    """校验扩展名与工具匹配（格式不符 → 可读错误，提示改用正确工具）。"""
    if target.suffix.lower() not in allowed:
        raise ToolError(
            f"{tool_name} 仅支持 {'/'.join(allowed)} 文件，收到：{target.name}"
            "（请确认素材格式并改用对应 read_* 工具）"
        )


def _cell_text(value: Any) -> str:
    """单元格值转字符串：空值保空（不破坏列对齐），转义表格分隔符。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def render_markdown_table(rows: list[list[str]]) -> str:
    """二维文本渲染为 Markdown 表格：列宽取最大列数，缺格补空以保持对齐。"""
    if not rows:
        return "（空表格）"
    width = max(len(row) for row in rows)
    normalized = [[row[i] if i < len(row) else "" for i in range(width)] for row in rows]
    lines = [
        "| " + " | ".join(normalized[0]) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return "\n".join(lines)


def _coerce_rows(value: Any, *, default: int, param: str = "max_rows") -> int:
    """解析行数参数（LLM 输入不可信：类型/范围校验，超硬上限截断）。"""
    if value is None:
        return default
    if isinstance(value, bool) or not isinstance(value, int):
        raise ToolError(f"参数 {param} 必须为整数")
    if value <= 0:
        raise ToolError(f"参数 {param} 必须为正整数")
    return min(value, _HARD_ROW_CAP)


def _decode_text(target: Path) -> tuple[str, str]:
    """按 utf-8 → gbk 顺序解码素材，返回 (文本, 实际编码)。"""
    data = target.read_bytes()
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8(容错替换)"


def _sniff_delimiter(text: str) -> str:
    """识别 CSV 分隔符（逗号 / 分号 / 制表符 / 竖线），失败退回逗号。"""
    sample = text[:4096]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _iter_block_items(document: Any) -> list[Any]:
    """按文档原顺序返回段落与表格（python-docx >= 1.1 提供 iter_inner_content）。"""
    iterator = getattr(document, "iter_inner_content", None)
    if callable(iterator):
        return list(iterator())
    from docx.oxml.ns import qn  # pragma: no cover - 旧版本兼容分支
    from docx.table import Table  # pragma: no cover
    from docx.text.paragraph import Paragraph  # pragma: no cover

    items: list[Any] = []  # pragma: no cover
    for child in document.element.body.iterchildren():  # pragma: no cover
        if child.tag == qn("w:p"):
            items.append(Paragraph(child, document))
        elif child.tag == qn("w:tbl"):
            items.append(Table(child, document))
    return items


def _heading_level(paragraph: Any) -> int | None:
    """段落标题层级：Title→1、Subtitle→2、Heading N→N，正文返回 None。"""
    style = getattr(paragraph, "style", None)
    name = str(getattr(style, "name", "") or "").strip()
    match = _HEADING_RE.match(name)
    if match:
        return max(1, min(6, int(match.group(1))))
    lowered = name.lower()
    if lowered == "title":
        return 1
    if lowered == "subtitle":
        return 2
    return None


class _ReadToolBase(OfficeTool):
    """只读素材工具基类：统一路径校验、扩展名校验、异常兜底与输出截断。"""

    #: 允许的扩展名（小写）；空元组表示不限制
    allowed_suffixes: tuple[str, ...] = ()

    def execute(self, args: dict[str, Any]) -> str:
        """读取素材并返回文本；失败抛 ``ToolError``（不抛未捕获异常）。"""
        filename = require_str(args, "filename", max_len=500)
        target = resolve_read_path(self.work_dir, filename)
        if self.allowed_suffixes:
            _require_suffix(target, self.allowed_suffixes, self.name)
        try:
            text = self.read_file(target, args)
        except ToolError:
            raise
        except Exception as exc:  # noqa: BLE001 — 解析异常转可读错误，不终止工具循环
            logger.warning(
                "tool_read_failed",
                extra={"event": "tool_read_failed", "tool": self.name, "error": str(exc)},
            )
            raise ToolError(
                f"{self.name} 读取失败（{target.name}）：{type(exc).__name__}: {exc}"
            ) from exc
        logger.info(
            "tool_read_done",
            extra={
                "event": "tool_read_done",
                "tool": self.name,
                "file": str(target),
                "chars": len(text),
            },
        )
        return truncate_output(text)

    @abstractmethod
    def read_file(self, target: Path, args: dict[str, Any]) -> str:
        """子类实现具体格式解析；``target`` 已通过工作目录越界校验且存在。"""


class ReadDocxTool(_ReadToolBase):
    """读取 .docx 素材：标题层级 / 段落文本 / 表格。"""

    name = "read_docx"
    description = (
        "读取工作目录内已有的 Word 文档（.docx）内容：按原顺序抽取标题层级、"
        "段落文本与表格（表格转 Markdown，保留行列对应）。"
        "任务附带 .docx 素材、需要基于素材内容改写/汇总/续写时，先用它读取，"
        "不要凭空编写内容。只读，不会修改文件。"
    )
    parameters: dict[str, Any] = {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "工作目录内的相对路径，如：素材/需求.docx"},
        },
        "required": ["filename"],
    }
    allowed_suffixes = (".docx", ".docm")

    def read_file(self, target: Path, args: dict[str, Any]) -> str:
        """解析 docx：标题→Markdown 标题、段落原文、表格→[[TABLE n]] + Markdown。"""
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(target))
        lines: list[str] = []
        paragraphs = 0
        tables = 0
        headings = 0
        for item in _iter_block_items(document):
            if isinstance(item, Paragraph):
                text = (item.text or "").strip()
                if not text:
                    continue
                level = _heading_level(item)
                if level is not None:
                    headings += 1
                    lines.append(f"{'#' * level} {text}")
                else:
                    paragraphs += 1
                    lines.append(text)
            elif isinstance(item, Table):
                tables += 1
                lines.append(f"[[TABLE {tables}]]")
                lines.append(
                    render_markdown_table(
                        [[_cell_text(cell.text) for cell in row.cells] for row in item.rows]
                    )
                )
        header = f"文件：{target.name}（{paragraphs} 段落 / {headings} 标题 / {tables} 表格）"
        return "\n".join([header, *lines]) if lines else f"{header}\n（文档无正文内容）"


class ReadXlsxTool(_ReadToolBase):
    """读取 .xlsx 素材：逐 sheet 返回，空单元格不破坏列对齐。"""

    name = "read_xlsx"
    description = (
        "读取工作目录内已有的 Excel 工作簿（.xlsx/.xlsm）：逐工作表返回单元格内容，"
        "空单元格保留占位以维持列对齐。任务附带表格数据、需要按数据填表/统计/汇总时，"
        "先用它读取；可用 sheet 指定单个工作表，max_rows 控制每表读取行数。只读，不修改文件。"
    )
    parameters: dict[str, Any] = {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "工作目录内的相对路径，如：数据/销量表.xlsx"},
            "sheet": {"type": "string", "description": "可选：只读取指定工作表名；缺省读取全部工作表"},
            "max_rows": {
                "type": "integer",
                "description": f"可选：每个工作表最多读取行数（默认 {DEFAULT_SHEET_ROWS}）",
            },
        },
        "required": ["filename"],
    }
    allowed_suffixes = (".xlsx", ".xlsm")

    def read_file(self, target: Path, args: dict[str, Any]) -> str:
        """按 sheet 遍历并渲染为 Markdown 表格（首行为表头）。"""
        from openpyxl import load_workbook

        sheet_arg = args.get("sheet")
        if sheet_arg is not None and not isinstance(sheet_arg, str):
            raise ToolError("参数 sheet 必须为字符串（工作表名）")
        max_rows = _coerce_rows(args.get("max_rows"), default=DEFAULT_SHEET_ROWS)

        workbook = load_workbook(str(target), read_only=True, data_only=True)
        try:
            names = list(workbook.sheetnames)
            if not names:
                return f"文件：{target.name}\n（工作簿无工作表）"
            wanted = (sheet_arg or "").strip()
            if wanted:
                if wanted not in names:
                    raise ToolError(f"工作表不存在：{wanted}（可选：{'、'.join(names)}）")
                names = [wanted]
            blocks = [f"文件：{target.name}（共 {len(workbook.sheetnames)} 个工作表）"]
            for name in names:
                sheet = workbook[name]
                rows: list[list[str]] = []
                truncated_rows = False
                for row in sheet.iter_rows(values_only=True):
                    if len(rows) >= max_rows:
                        truncated_rows = True
                        break
                    rows.append([_cell_text(value) for value in row])
                # 去除尾部全空行（read_only 模式对空工作表会返回一个空行占位）
                while rows and not any(cell for cell in rows[-1]):
                    rows.pop()
                blocks.append(f"[[SHEET {name}]]")
                if not rows:
                    blocks.append("（空工作表）")
                    continue
                blocks.append(render_markdown_table(rows))
                if truncated_rows:
                    blocks.append(f"（仅显示前 {max_rows} 行，可用 max_rows 调整）")
            return "\n".join(blocks)
        finally:
            workbook.close()


class ReadCsvTool(_ReadToolBase):
    """读取 .csv/.tsv 素材：表头 + 前 N 行，编码容错。"""

    name = "read_csv"
    description = (
        "读取工作目录内的 CSV/TSV 表格文件：返回表头与前 N 行数据（默认 50 行，"
        "max_rows 可调），分隔符自动识别，编码按 utf-8 → gbk 兜底。"
        "任务附带 .csv 素材时先用它读取，再据此产出；不要凭猜测编写表内数据。只读，不修改文件。"
    )
    parameters: dict[str, Any] = {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "工作目录内的相对路径，如：清单.csv"},
            "max_rows": {
                "type": "integer",
                "description": f"可选：最多返回的数据行数（不含表头，默认 {DEFAULT_CSV_ROWS}）",
            },
        },
        "required": ["filename"],
    }
    allowed_suffixes = (".csv", ".tsv")

    def read_file(self, target: Path, args: dict[str, Any]) -> str:
        """解析 CSV：表头 + 前 N 行，返回文本标注编码与行数。"""
        max_rows = _coerce_rows(args.get("max_rows"), default=DEFAULT_CSV_ROWS)
        text, encoding = _decode_text(target)
        delimiter = "\t" if target.suffix.lower() == ".tsv" else _sniff_delimiter(text)
        rows: list[list[str]] = []
        total = 0
        for row in csv.reader(io.StringIO(text), delimiter=delimiter):
            total += 1
            if len(rows) <= max_rows:
                rows.append([_cell_text(cell) for cell in row])
        data_rows = max(total - 1, 0)
        header = (
            f"文件：{target.name}（编码 {encoding}，分隔符 {delimiter!r}，"
            f"共 {data_rows} 行数据）"
        )
        if not rows:
            return f"{header}\n（文件为空）"
        body = render_markdown_table(rows)
        if data_rows > max_rows:
            body += f"\n（仅显示前 {max_rows} 行数据，可用 max_rows 调整）"
        return f"{header}\n{body}"


class ReadTextTool(_ReadToolBase):
    """读取纯文本素材（.txt/.md/.json/.log 等）：超上限截断并标注。"""

    name = "read_text"
    description = (
        "读取工作目录内的纯文本素材（.txt/.md/.json/.log/.yaml 等），按文本原样返回，"
        f"超过 {MAX_OUTPUT_CHARS} 字符自动截断并标注 truncated。"
        "任务附带文本类素材、需要基于原文产出（改写/提取要点/格式转换）时先用它读取；"
        "二进制或 Office 文件请改用 read_docx / read_xlsx / read_csv。只读，不修改文件。"
    )
    parameters: dict[str, Any] = {
        "type": "object",
        "properties": {
            "filename": {"type": "string", "description": "工作目录内的相对路径，如：素材/背景.md"},
        },
        "required": ["filename"],
    }
    allowed_suffixes = _TEXT_EXTENSIONS

    def read_file(self, target: Path, args: dict[str, Any]) -> str:
        """按编码兜底解码纯文本，返回带文件信息的原文（超长由基类截断）。"""
        text, encoding = _decode_text(target)
        header = f"文件：{target.name}（编码 {encoding}，{len(text)} 字符）"
        return f"{header}\n{text}" if text else f"{header}\n（文件为空）"
