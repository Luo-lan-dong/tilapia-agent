"""产出增强单元测试（ADR-065 M3：UT-M3-1 / UT-M3-2 / UT-M3-3）。

三个用例族共同的核心红线是「**默认参数行为不变**」：不给增强参数时，
write_docx / write_pptx 的产物结构与 M2 完全一致（不引入域、不引入图表部件）。
"""

import zipfile
from pathlib import Path
from typing import Any

import pytest

from app.tools.base import ToolError, clean_llm_text
from app.tools.chart_tool import RenderChartTool
from app.tools.docx_tool import DocxTool
from app.tools.pptx_tool import PptxTool

#: 目录域代码关键字（Word 域指令以 instrText 形式写入 document.xml）
_TOC_FIELD = "TOC"
#: 页码域代码关键字（写入 footer*.xml）
_PAGE_FIELD = "PAGE"


def write_docx(work_dir: Path, **extra: Any) -> str:
    """调用 write_docx 渲染固定两段内容，返回工具结果文本。"""
    args: dict[str, Any] = {
        "filename": "报告.docx",
        "title": "季度报告",
        "blocks": [
            {"type": "heading", "text": "第一章 概览", "level": 1},
            {"type": "paragraph", "text": "正文内容。"},
        ],
    }
    args.update(extra)
    return DocxTool(work_dir).execute(args)


def write_pptx(work_dir: Path, **extra: Any) -> str:
    """调用 write_pptx 渲染固定两页内容，返回工具结果文本。"""
    args: dict[str, Any] = {
        "filename": "汇报.pptx",
        "title": "季度汇报",
        "slides": [
            {"title": "概览", "bullets": ["要点一", "要点二"]},
            {"title": "数据", "bullets": ["要点三"]},
        ],
    }
    args.update(extra)
    return PptxTool(work_dir).execute(args)


def part_texts(path: Path) -> dict[str, str]:
    """解压 Office 包并返回「部件名 → UTF-8 文本」（docx/pptx 同为 OPC 包）。"""
    with zipfile.ZipFile(path) as archive:
        return {
            name: archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.endswith(".xml")
        }


def part_names(path: Path, prefix: str) -> list[str]:
    """返回包内指定前缀的部件名（如 word/media/、ppt/media/、ppt/charts/）。"""
    with zipfile.ZipFile(path) as archive:
        return [name for name in archive.namelist() if name.startswith(prefix)]


def footer_parts(path: Path) -> dict[str, str]:
    """取全部页脚部件（footer1.xml 等）文本。"""
    return {
        name: text
        for name, text in part_texts(path).items()
        if Path(name).name.startswith("footer")
    }


def header_parts(path: Path) -> dict[str, str]:
    """取全部页眉部件（header1.xml 等）文本。"""
    return {
        name: text
        for name, text in part_texts(path).items()
        if Path(name).name.startswith("header")
    }


class TestDocxEnhancement:
    """UT-M3-1：目录域 / 页码域 / 页眉页脚写入，且默认参数不引入域。"""

    def test_toc_fields_and_header_footer_written(self, tmp_path: Path) -> None:
        result = write_docx(
            tmp_path,
            toc=True,
            page_numbers=True,
            header="机密 · 内部资料",
            footer="示例公司",
        )
        target = tmp_path / "报告.docx"
        assert target.is_file()

        # 目录域：document.xml 中必须出现 TOC 域代码与 fldChar 结构
        document_xml = part_texts(target)["word/document.xml"]
        assert _TOC_FIELD in document_xml
        assert 'TOC \\o "1-3" \\h \\z \\u' in document_xml
        assert "fldChar" in document_xml
        # 目录位于主标题之后、正文之前（插在 body 位置 1）
        assert document_xml.index("目录") < document_xml.index("第一章 概览")
        # updateFields 标记：提示 Word/WPS 打开时更新域
        assert "updateFields" in part_texts(target)["word/settings.xml"]

        # 页脚：页码域与页脚文本共存（页码另起一行）
        footers = footer_parts(target)
        assert footers, "增强模式下必须生成页脚部件"
        footer_xml = "".join(footers.values())
        assert _PAGE_FIELD in footer_xml
        assert "fldChar" in footer_xml
        assert "示例公司" in footer_xml
        assert footer_xml.count("<w:p>") + footer_xml.count("<w:p ") >= 2

        # 页眉：文本落位
        assert "机密 · 内部资料" in "".join(header_parts(target).values())

        # 返回文本必须说明「域需打开文档更新」（避免误判为失败）
        assert "域" in result
        assert "F9" in result
        assert "报告.docx" in result

    def test_toc_only_and_page_numbers_only(self, tmp_path: Path) -> None:
        # 仅目录：不生成页脚部件
        write_docx(tmp_path, filename="仅目录.docx", toc=True)
        only_toc = part_texts(tmp_path / "仅目录.docx")
        assert _TOC_FIELD in only_toc["word/document.xml"]
        assert not [name for name in only_toc if Path(name).name.startswith("footer")]

        # 仅页码：生成页脚部件但正文无目录域
        write_docx(tmp_path, filename="仅页码.docx", page_numbers=True)
        only_page = part_texts(tmp_path / "仅页码.docx")
        assert _PAGE_FIELD in "".join(
            text for name, text in only_page.items() if Path(name).name.startswith("footer")
        )
        assert "fldChar" not in only_page["word/document.xml"]

    def test_page_numbers_without_footer_text_single_paragraph(self, tmp_path: Path) -> None:
        write_docx(tmp_path, filename="页码.docx", page_numbers=True)
        footer_xml = "".join(footer_parts(tmp_path / "页码.docx").values())
        assert _PAGE_FIELD in footer_xml
        # 无页脚文本时页码直接落在首个段落（不多出空段落）
        assert footer_xml.count("<w:p>") + footer_xml.count("<w:p ") == 1
        assert "w:jc" in footer_xml  # 居中

    def test_toc_position_with_table_first_block(self, tmp_path: Path) -> None:
        """首块为表格时目录仍插在标题之后、正文之前（位置不依赖块类型）。"""
        from docx import Document

        DocxTool(tmp_path).execute({
            "filename": "带表.docx",
            "title": "带表报告",
            "blocks": [
                {"type": "table", "rows": [["指标", "值"], ["预算", "100"]]},
                {"type": "heading", "text": "第二章", "level": 1},
            ],
            "toc": True,
        })
        target = tmp_path / "带表.docx"
        # python-docx 可正常重开（XML 合法）
        document = Document(str(target))
        elements = list(document.element.body)
        # body = [Title, 目录标题, 目录域, 表格, Heading1]：目录（下标 1/2）必须在表格（下标 3）之前
        kinds = [element.tag.rsplit("}", 1)[-1] for element in elements]
        assert kinds[0] == "p"
        assert kinds[1] == "p" and kinds[2] == "p"
        assert kinds.index("tbl") == 3
        texts = ["".join(element.itertext()) for element in elements]
        assert "目录" in texts[1]
        assert "TOC" in texts[2]
        assert "TOC" in part_texts(target)["word/document.xml"]

    def test_default_output_has_no_fields_or_parts(self, tmp_path: Path) -> None:
        """默认参数：产物结构与改动前一致（无域、无页眉页脚部件）。"""
        result = write_docx(tmp_path)
        target = tmp_path / "报告.docx"
        parts = part_texts(target)
        assert "fldChar" not in parts["word/document.xml"]
        assert _TOC_FIELD not in parts["word/document.xml"]
        assert "updateFields" not in parts["word/settings.xml"]
        assert not header_parts(target)
        assert not footer_parts(target)
        # 返回文本不出现域更新提示（未附加任何增强）
        assert "域" not in result
        # 正文内容照常写入
        from docx import Document

        document = Document(str(target))
        texts = [paragraph.text for paragraph in document.paragraphs]
        assert texts[0] == "季度报告"
        assert "第一章 概览" in texts
        assert "正文内容。" in texts

    def test_invalid_enhancement_args_readable_error(self, tmp_path: Path) -> None:
        with pytest.raises(ToolError) as excinfo:
            write_docx(tmp_path, toc="yes")
        assert "toc" in str(excinfo.value)
        with pytest.raises(ToolError) as excinfo:
            write_docx(tmp_path, page_numbers=1)
        assert "page_numbers" in str(excinfo.value)
        with pytest.raises(ToolError) as excinfo:
            write_docx(tmp_path, header=123)
        assert "header" in str(excinfo.value)
        with pytest.raises(ToolError) as excinfo:
            write_docx(tmp_path, footer=["页脚"])
        assert "footer" in str(excinfo.value)
        # 非法输入不产生文件
        assert not (tmp_path / "报告.docx").exists()


class TestPptxEnhancement:
    """UT-M3-2：原生图表 + 演讲者备注写入，且默认参数不新增图表。"""

    def test_charts_and_notes_written(self, tmp_path: Path) -> None:
        from pptx import Presentation

        result = write_pptx(
            tmp_path,
            charts=[
                {
                    "slide": 2,
                    "type": "bar",
                    "categories": ["一季度", "二季度"],
                    "series": [{"name": "销售额", "values": [10, 20]}],
                    "title": "季度销售",
                },
                {
                    "slide": 3,
                    "type": "line",
                    "categories": ["一月", "二月", "三月"],
                    "series": [
                        {"name": "收入", "values": [1, 2, 3]},
                        {"name": "成本", "values": [1, 1, 2]},
                    ],
                },
            ],
            notes=[{"slide": 1, "text": "开场白：介绍项目背景"}, {"slide": 3, "text": "强调趋势"}],
        )
        target = tmp_path / "汇报.pptx"
        assert target.is_file()
        assert "2 张原生图表" in result
        assert "2 页演讲者备注" in result

        presentation = Presentation(str(target))
        slides = list(presentation.slides)
        assert len(slides) == 3  # 封面 + 2 页内容

        # 封面只写备注，不含图表（NFKC 会把全角冒号归一为半角，故比对归一后文本）
        assert not [shape for shape in slides[0].shapes if shape.has_chart]
        assert clean_llm_text("开场白：介绍项目背景") in slides[0].notes_slide.notes_text_frame.text

        bar_chart = next(shape for shape in slides[1].shapes if shape.has_chart).chart
        assert str(bar_chart.chart_type).startswith("COLUMN_CLUSTERED")
        assert [series.name for series in bar_chart.plots[0].series] == ["销售额"]
        assert list(bar_chart.plots[0].series[0].values) == [10.0, 20.0]
        assert list(bar_chart.plots[0].categories) == ["一季度", "二季度"]
        assert bar_chart.has_title is True
        assert bar_chart.chart_title.text_frame.text == "季度销售"

        line_chart = next(shape for shape in slides[2].shapes if shape.has_chart).chart
        assert str(line_chart.chart_type).startswith("LINE")
        assert [series.name for series in line_chart.plots[0].series] == ["收入", "成本"]
        assert list(line_chart.plots[0].series[1].values) == [1.0, 1.0, 2.0]
        assert line_chart.has_title is False
        assert "强调趋势" in slides[2].notes_slide.notes_text_frame.text

        # 图表必须是真正的 chart 部件（可在 PowerPoint 中编辑数据）
        chart_parts = part_names(target, "ppt/charts/chart")
        assert len(chart_parts) == 2

    def test_pie_chart_uses_first_series(self, tmp_path: Path) -> None:
        from pptx import Presentation

        write_pptx(
            tmp_path,
            filename="饼图.pptx",
            charts=[
                {
                    "slide": 2,
                    "type": "pie",
                    "categories": ["甲", "乙", "丙"],
                    "series": [{"name": "占比", "values": [5, 3, 2]}],
                }
            ],
        )
        presentation = Presentation(str(tmp_path / "饼图.pptx"))
        chart = next(
            shape.chart for shape in list(presentation.slides)[1].shapes if shape.has_chart
        )
        assert str(chart.chart_type).startswith("PIE")
        assert list(chart.plots[0].series[0].values) == [5.0, 3.0, 2.0]
        assert list(chart.plots[0].categories) == ["甲", "乙", "丙"]

    def test_default_output_has_no_charts(self, tmp_path: Path) -> None:
        """默认参数：不新增图表部件，产物结构与改动前一致。"""
        from pptx import Presentation

        result = write_pptx(tmp_path)
        target = tmp_path / "汇报.pptx"
        presentation = Presentation(str(target))
        assert not [
            shape for slide in presentation.slides for shape in slide.shapes if shape.has_chart
        ]
        assert not part_names(target, "ppt/charts/")
        assert "图表" not in result
        assert "备注" not in result
        assert "汇报.pptx" in result

    def test_notes_per_slide_still_supported(self, tmp_path: Path) -> None:
        """slides[].notes（M2 既有能力）与 notes 参数（M3）并存不冲突（同页追加不覆盖）。"""
        from pptx import Presentation

        write_pptx(
            tmp_path,
            filename="备注.pptx",
            slides=[
                {"title": "概览", "bullets": ["一"], "notes": "页内备注"},
                {"title": "数据", "bullets": ["二"]},
            ],
            notes=[{"slide": 2, "text": "外部补充备注"}],
        )
        presentation = Presentation(str(tmp_path / "备注.pptx"))
        slides = list(presentation.slides)
        # notes 页码 1 起算且含封面：slide=2 → 页序第 2 张（即第一张内容页，list 下标 1）
        second = slides[1].notes_slide.notes_text_frame.text
        assert "页内备注" in second  # M2 的 slides[].notes 保留
        assert "外部补充备注" in second  # M3 的 notes 同页追加，不覆盖原有备注
        # 未指定的页不被写入
        assert slides[2].notes_slide.notes_text_frame.text == ""

    def test_invalid_charts_and_notes_readable_error(self, tmp_path: Path) -> None:
        base_chart: dict[str, Any] = {
            "slide": 2,
            "type": "bar",
            "categories": ["一", "二"],
            "series": [{"name": "销售额", "values": [1, 2]}],
        }
        cases: list[dict[str, Any]] = [
            {"charts": [dict(base_chart, type="scatter")]},
            {"charts": [dict(base_chart, slide=9)]},
            {"charts": [dict(base_chart, slide="2")]},
            {"charts": [dict(base_chart, categories=[])]},
            {"charts": [dict(base_chart, series=[])]},
            {"charts": [dict(base_chart, series=[{"name": "s", "values": [1]}])]},
            {"charts": [dict(base_chart, series=[{"name": "", "values": [1, 2]}])]},
            {"charts": [dict(base_chart, series=[{"name": "s", "values": ["x", "y"]}])]},
            {"charts": "bar"},
            {"notes": [{"slide": 9, "text": "越界"}]},
            {"notes": [{"slide": 2, "text": "  "}]},
        ]
        for extra in cases:
            with pytest.raises(ToolError) as excinfo:
                write_pptx(tmp_path, filename="坏.pptx", **extra)
            message = str(excinfo.value)
            assert message.strip()
            assert "Traceback" not in message
        # 全部非法输入均在落盘前拦截
        assert not (tmp_path / "坏.pptx").exists()


class TestRenderChart:
    """UT-M3-3：render_chart 产出可用 PNG（PIL 可读、可嵌入 docx/pptx）。"""

    def test_png_render_and_embed_into_docx(self, tmp_path: Path) -> None:
        from PIL import Image

        result = RenderChartTool(tmp_path).execute({
            "filename": "图表/季度销售.png",
            "type": "bar",
            "categories": ["一季度", "二季度", "三季度"],
            "series": [
                {"name": "销售额", "values": [12, 20, 33]},
                {"name": "成本", "values": [8, 11, 15]},
            ],
            "title": "季度销售趋势（中文标题）",
            "xlabel": "季度",
            "ylabel": "万元",
        })
        target = tmp_path / "图表" / "季度销售.png"
        assert target.is_file()
        # 返回文本含实际路径与尺寸
        assert "季度销售.png" in result
        assert "8×4.5 英寸" in result
        assert "150dpi" in result
        assert "1200×675 像素" in result
        assert "2 个系列 × 3 个数据点" in result

        with Image.open(str(target)) as image:
            assert image.format == "PNG"
            assert image.size == (1200, 675)

        # 可被 python-docx 以图片方式嵌入既有文档
        from docx import Document
        from docx.shared import Inches

        document = Document()
        document.add_heading("配图报告", level=1)
        document.add_picture(str(target), width=Inches(5))
        docx_path = tmp_path / "配图.docx"
        document.save(str(docx_path))
        images = part_names(docx_path, "word/media/")
        assert images, "docx 内应包含嵌入图片部件"
        assert Document(str(docx_path)).inline_shapes[0].width > 0

    def test_line_and_pie_variants(self, tmp_path: Path) -> None:
        from PIL import Image

        tool = RenderChartTool(tmp_path)
        line = tool.execute({
            "filename": "line.png",
            "type": "line",
            "categories": ["一月", "二月"],
            "series": [{"name": "收入", "values": [1, 2]}],
            "width": 6,
            "height": 3,
        })
        assert "6×3 英寸" in line
        assert "900×450 像素" in line
        pie = tool.execute({
            "filename": "pie",
            "type": "pie",
            "categories": ["甲", "乙"],
            "series": [{"name": "占比", "values": [3, 7]}],
            "title": "占比分布",
        })
        assert (tmp_path / "pie.png").is_file()  # 缺扩展名自动补 .png
        with Image.open(str(tmp_path / "pie.png")) as image:
            assert image.size == (1200, 675)
        assert "pie.png" in pie

    def test_png_can_be_embedded_into_pptx(self, tmp_path: Path) -> None:
        """PNG 可经 python-pptx 以图片方式插入幻灯片（产出增强闭环）。"""
        from pptx import Presentation
        from pptx.util import Inches

        RenderChartTool(tmp_path).execute({
            "filename": "素材.png",
            "type": "bar",
            "categories": ["甲", "乙"],
            "series": [{"name": "数量", "values": [3, 5]}],
        })
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.add_picture(str(tmp_path / "素材.png"), Inches(1), Inches(1), width=Inches(6))
        deck = tmp_path / "带图.pptx"
        presentation.save(str(deck))
        assert part_names(deck, "ppt/media/")

    def test_invalid_args_readable_error(self, tmp_path: Path) -> None:
        tool = RenderChartTool(tmp_path)
        valid: dict[str, Any] = {
            "filename": "c.png",
            "type": "bar",
            "categories": ["甲", "乙"],
            "series": [{"name": "数量", "values": [1, 2]}],
        }
        cases: list[dict[str, Any]] = [
            {"type": "scatter"},
            {"type": ""},
            {"series": []},
            {"series": [{"name": "数量", "values": [1]}]},
            {"series": [{"name": "", "values": [1, 2]}]},
            {"series": [{"name": "数量", "values": ["一", "二"]}]},
            {"series": ["不是对象"]},
            {"categories": []},
            {"categories": [None, "乙"]},
            {"filename": "../越界.png"},
            {"filename": "  "},
            {"width": "8"},
            {"height": 0.1},
            {"title": 3},
        ]
        for override in cases:
            args = {**valid, **override}
            with pytest.raises(ToolError) as excinfo:
                tool.execute(args)
            assert str(excinfo.value).strip()
            assert "Traceback" not in str(excinfo.value)
        assert not (tmp_path / "c.png").exists()

    def test_chinese_title_does_not_raise(self, tmp_path: Path) -> None:
        """中文标题/标签必须可渲染（字体兜底；不得因缺字形抛异常）。"""
        tool = RenderChartTool(tmp_path)
        result = tool.execute({
            "filename": "中文.png",
            "type": "line",
            "categories": ["华东区", "华南区"],
            "series": [{"name": "同比增长率（%）", "values": [-3, 8]}],
            "title": "各区同比增长（负值测试 −）",
            "xlabel": "区域",
            "ylabel": "百分比",
        })
        assert (tmp_path / "中文.png").is_file()
        assert "中文.png" in result


class TestRegistry:
    """注册表：render_chart 已注册（共 12 个工具）。"""

    def test_render_chart_registered(self, tmp_path: Path) -> None:
        import os

        from app.config import AgentConfig
        from app.executor.tool_loop import default_registry

        os.environ.setdefault("AGENT_REGISTRATION_TOKEN", "t")
        registry = default_registry(tmp_path, AgentConfig(_env_file=None))
        names = [schema["function"]["name"] for schema in registry.openai_schemas()]
        assert "render_chart" in names
        assert len(names) == 12
        assert names[-1] == "render_chart"
        # schema 自证可被 LLM 调用（必填参数齐备）
        schema = registry.get("render_chart").parameters
        assert schema["required"] == ["filename", "type", "categories", "series"]
        assert schema["properties"]["type"]["enum"] == ["line", "bar", "pie"]

    def test_chart_png_roundtrip_via_io(self, tmp_path: Path) -> None:
        """PNG 头部自证为合法图片（不依赖 PIL：解析 IHDR 宽高）。"""
        RenderChartTool(tmp_path).execute({
            "filename": "raw.png",
            "type": "bar",
            "categories": ["甲"],
            "series": [{"name": "数量", "values": [1]}],
        })
        data = (tmp_path / "raw.png").read_bytes()
        assert data[:8] == b"\x89PNG\r\n\x1a\n"
        # PNG IHDR：宽度、高度均为 4 字节大端整数（偏移 16 / 20）
        assert int.from_bytes(data[16:20], "big") == 1200
        assert int.from_bytes(data[20:24], "big") == 675
