"""交付质量闸测试（ADR-065 M4：UT-M4-1/2/3/4 + L2 证据位置判据）。

沿用 ``test_office_bridge`` 的 fixture 范式（``make_config`` / ``FakeCallbacks`` /
monkeypatch 桥接层 ``LLMClient`` 注入脚本化 LLM），交付物一律用真实库写出并真打开，
不做结构性 mock。
"""

import asyncio
import hashlib
import os
import time
from contextlib import contextmanager
from pathlib import Path

import pytest
from docx import Document

from app.config import AgentConfig
from app.executor.llm_client import LLMError
from app.executor.office_bridge import BridgeCallbacks, OfficeBridge
from app.executor.verify import (
    MAX_REPAIR_ROUNDS,
    VerifyPlan,
    run_verify,
    run_verify_loop,
)
from app.models.schemas import TaskDispatchPayload, TaskSpecification
from app.tools.base import clean_llm_text

#: 通过判据所需的「已达标」段落（回灌后必须逐字保留）
PASSING_PARAGRAPHS = ("市场规模 2024 年为 12 亿元。", "竞争格局：头部三家合计占据 60% 份额。")
#: 缺失章节（首轮不写，回灌轮才补）
MISSING_HEADING = "风险提示"
MISSING_PARAGRAPH = "风险提示：原材料价格波动可能影响毛利。"

#: 可被 L1 自动判定的 checklist（关键词条目；不含语义条目，避免 L2 干扰）
CHECKLIST = "- [ ] 报告必须包含「市场概览」章节\n- [ ] 报告必须包含「风险提示」章节\n"

#: 含 L2 语义条目的 checklist
SEMANTIC_CHECKLIST = "- [ ] 结论需给出可执行的建议\n"


def make_config(**overrides: str) -> AgentConfig:
    values = {
        "AGENT_REGISTRATION_TOKEN": "t",
        "WORKSPACE_ROOT": "",
        "OFFICE_LLM_MAX_ROUNDS": "10",
    }
    values.update({k: str(v) for k, v in overrides.items()})
    for key, value in values.items():
        os.environ[key] = value
    return AgentConfig(_env_file=None)


def make_dispatch(**extra) -> TaskDispatchPayload:
    payload = {
        "task_id": "task-verify-001",
        "task_name": "写调研报告",
        "description": "调研某主题并产出 docx",
        "executor_model": {
            "source": "hub",
            "model": "deepseek/deepseek-chat",
            "provider": "deepseek",
            "api_key": "sk-test-123",
            "base_url": "https://api.deepseek.com",
        },
        "output_dir": str(extra.pop("output_dir")) if "output_dir" in extra else None,
    }
    payload.update(extra)
    return TaskDispatchPayload.model_validate(payload)


class FakeCallbacks(BridgeCallbacks):
    """收集全部上报事件的假回调（含验收阻断，ADR-065 M4）。"""

    def __init__(self) -> None:
        self.events: list[tuple] = []

    async def on_progress(
        self, task_id, progress, todos=None, *, stage=None, message=None, tool=None, detail=None
    ):
        self.events.append(("progress", task_id, progress, stage, message, todos))

    async def on_task_started(self, task_id, executor_model_actual):
        self.events.append(("started", task_id, executor_model_actual))

    async def on_completed(
        self, task_id, session_id, plan_content=None, change_summary=None, project_context=None
    ):
        self.events.append(("completed", task_id, session_id, change_summary))

    async def on_failed(self, task_id, reason):
        self.events.append(("failed", task_id, reason))

    async def on_file_tree(self, task_id, files, truncated):
        self.events.append(("file_tree", task_id, [item["path"] for item in files], truncated))

    async def on_file_chunk(self, request_id, seq, total, data_b64, done):
        self.events.append(("file_chunk", request_id, seq, total, done))

    async def on_file_error(self, request_id, reason):
        self.events.append(("file_error", request_id, reason))

    async def on_verify_blocked(self, task_id, session_id, report):
        self.events.append(("verify_blocked", task_id, session_id, report))

    def kinds(self) -> list[str]:
        """事件类型序列。"""
        return [event[0] for event in self.events]

    def of(self, kind: str) -> list[tuple]:
        """按类型取事件。"""
        return [event for event in self.events if event[0] == kind]


def write_docx(path: Path, title: str, paragraphs: list[str], headings: list[str]) -> None:
    """用 python-docx 真实写出交付物（文档标题 + 章节标题 + 正文段落）。"""
    document = Document()
    document.add_heading(title, level=0)
    for heading in headings:
        document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(str(path))


def docx_paragraphs(path: Path) -> list[str]:
    """读取 docx 全部非空段落文本（回灌前后逐字比对用）。"""
    return [p.text for p in Document(str(path)).paragraphs if p.text.strip()]


def digest(text: str) -> str:
    """段落文本 sha256（UT-M4-3 断言已达标内容未被重写）。"""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def rate_docx(work_dir: Path, *, headings: list[str], paragraphs: list[str] | None = None) -> Path:
    """写出「报告.docx」交付物并返回路径。"""
    target = work_dir / "报告.docx"
    write_docx(target, "市场调研报告", list(paragraphs or PASSING_PARAGRAPHS), headings)
    return target


def tool_call_response(call_id: str, blocks: list[dict]) -> dict:
    """构造 write_docx 工具调用的 LLM 响应。"""
    import json

    return {
        "message": {
            "role": "assistant",
            "content": "",
            "tool_calls": [
                {
                    "id": call_id,
                    "type": "function",
                    "function": {
                        "name": "write_docx",
                        "arguments": json.dumps(
                            {
                                "filename": "报告.docx",
                                "title": "市场调研报告",
                                "blocks": blocks,
                            },
                            ensure_ascii=False,
                        ),
                    },
                }
            ],
        },
        "prompt_tokens": 10,
        "completion_tokens": 5,
    }


def final_response(text: str = "交付完成。") -> dict:
    """构造最终答复（不带工具调用）的 LLM 响应。"""
    return {
        "message": {"role": "assistant", "content": text},
        "prompt_tokens": 10,
        "completion_tokens": 5,
    }


def initial_blocks() -> list[dict]:
    """首轮产出：只有「市场概览」章节（缺「风险提示」）。"""
    return [
        {"type": "heading", "text": MISSING_HEADING.replace(MISSING_HEADING, "市场概览"), "level": 1},
        *({"type": "paragraph", "text": text} for text in PASSING_PARAGRAPHS),
    ]


def repaired_blocks() -> list[dict]:
    """回灌轮产出：**原样保留已达标段落** + 补「风险提示」章节。"""
    return [
        {"type": "heading", "text": "市场概览", "level": 1},
        *({"type": "paragraph", "text": text} for text in PASSING_PARAGRAPHS),
        {"type": "heading", "text": MISSING_HEADING, "level": 1},
        {"type": "paragraph", "text": MISSING_PARAGRAPH},
    ]


class ScriptedVerifyLLM:
    """脚本化 LLM：首轮写缺章节的交付物；回灌轮按 ``fix_repair`` 决定是否补齐。

    ``fix_repair=False`` 用于「连续 5 轮不过 → verify_blocked」场景（模型不修正）。
    """

    model = "fake-model"

    def __init__(self, work_dir: Path, *, fix_repair: bool = True) -> None:
        self.work_dir = work_dir
        self.fix_repair = fix_repair
        self.calls = 0
        self.system_prompts: list[str] = []
        self.repair_prompts: list[str] = []
        #: 回灌前交付物段落文本摘要（断言已达标段落未被重写）
        self.preexisting_digests: list[str] = []
        self._written = False

    async def chat(self, messages, tools=None):  # noqa: ANN001, ANN202
        self.calls += 1
        self.system_prompts.append(str(messages[0].get("content") or ""))
        instruction = str(messages[1].get("content") or "") if len(messages) > 1 else ""
        repair = instruction.startswith("自动交付质量闸未通过")
        if repair and instruction not in self.repair_prompts:
            self.repair_prompts.append(instruction)
            target = self.work_dir / "报告.docx"
            if target.is_file():
                self.preexisting_digests = [digest(text) for text in docx_paragraphs(target)]
        if self._written:
            self._written = False
            return final_response()
        self._written = True
        blocks = repaired_blocks() if (repair and self.fix_repair) else initial_blocks()
        return tool_call_response(f"call-{self.calls}", blocks)


class JudgeLLM:
    """假 L2 判分模型：按预设内容返回（可注入异常以模拟判分不可用）。"""

    model = "judge-model"

    def __init__(self, payload: str | Exception) -> None:
        self.payload = payload
        self.calls = 0
        self.prompts: list[str] = []

    async def chat(self, messages, tools=None):  # noqa: ANN001, ANN202
        self.calls += 1
        self.prompts.append(str(messages[-1].get("content") or ""))
        if isinstance(self.payload, Exception):
            raise self.payload
        return {
            "message": {"role": "assistant", "content": self.payload},
            "prompt_tokens": 5,
            "completion_tokens": 5,
        }


@contextmanager
def patched_llm(fake):
    """把桥接层 LLMClient 换成脚本化假件（沿用既有测试的注入方式）。"""
    import app.executor.office_bridge as bridge_mod

    class FakeLLMClient:
        def __init__(self, **kwargs) -> None:
            self.model = kwargs.get("model", "fake-model")
            self._llm = fake

        async def chat(self, messages, tools=None):  # noqa: ANN001, ANN202
            return await self._llm.chat(messages, tools)

    original = bridge_mod.LLMClient
    bridge_mod.LLMClient = FakeLLMClient
    try:
        yield
    finally:
        bridge_mod.LLMClient = original


def plan_for(checklist: str) -> VerifyPlan:
    """构造只来自中枢 checklist 的判据计划。"""
    return VerifyPlan.from_text(checklist)


# --------------------------------------------------------------------------- #
# UT-M4-1 结构闸：空文件 / 损坏文件
# --------------------------------------------------------------------------- #


async def test_ut_m4_1_empty_docx_fails_l0(tmp_path: Path) -> None:
    """UT-M4-1：空文件 → L0 判不过，且给出条目级原因（不进入 L1）。"""
    plan = VerifyPlan.from_text("- [ ] 交付物「报告.docx」须包含「市场概览」章节")
    assert plan.deliverables == ["报告.docx"]
    (tmp_path / "报告.docx").write_bytes(b"")

    report = await run_verify(tmp_path, plan, since=0)

    assert report.passed is False
    assert report.layer == "L0"
    failed = {item.id: item for item in report.failures}
    assert "L0-D1-OPEN" in failed and "无法打开" in failed["L0-D1-OPEN"].reason
    assert "L0-D1-NONEMPTY" in failed and failed["L0-D1-NONEMPTY"].reason
    assert all(item.reason for item in report.failures)  # 每条都带原因
    assert all(item.layer == "L0" for item in report.items)  # L0 不过即短路，不跑 L1


async def test_ut_m4_1_broken_docx_fails_l0(tmp_path: Path) -> None:
    """UT-M4-1：损坏 docx（非 zip 字节）→ L0 判不过并指明解析库。"""
    plan = VerifyPlan.from_text("- [ ] 交付物「报告.docx」须包含「市场概览」章节")
    (tmp_path / "报告.docx").write_bytes(b"this is not a real docx package")

    report = await run_verify(tmp_path, plan, since=0)

    assert report.passed is False and report.layer == "L0"
    opened = next(item for item in report.items if item.id == "L0-D1-OPEN")
    assert opened.passed is False
    assert "python-docx" in opened.requirement
    assert "无法打开" in opened.reason


# --------------------------------------------------------------------------- #
# UT-M4-2 内容闸：缺章节 + 占位符
# --------------------------------------------------------------------------- #


async def test_ut_m4_2_missing_section_and_todo_fails_l1(tmp_path: Path) -> None:
    """UT-M4-2：缺关键章节 + 含 TODO 占位 → L1 判不过并指出缺失项。"""
    plan = plan_for(CHECKLIST)
    target = tmp_path / "报告.docx"
    write_docx(
        target,
        "市场调研报告",
        ["市场规模 2024 年为 12 亿元。", "结论：TODO 待补充。"],
        ["市场概览"],
    )

    report = await run_verify(tmp_path, plan, since=0)

    assert report.passed is False and report.layer == "L1"
    items = {item.id: item for item in report.items}
    # 已达标项通过（不参与回灌）
    assert items["L1-C1"].passed is True
    # 缺失章节被指出
    assert items["L1-C2"].passed is False
    assert "风险提示" in items["L1-C2"].reason and "缺失" in items["L1-C2"].reason
    # 占位符被指出（含命中位置）
    assert items["L1-PLACEHOLDER"].passed is False
    assert "TODO" in items["L1-PLACEHOLDER"].reason
    assert "第 4 段" in items["L1-PLACEHOLDER"].reason
    assert report.to_dict()["layer"] == "L1"


async def test_l1_empty_section_fails(tmp_path: Path) -> None:
    """L1 禁空章节：标题下没有任何正文 → 判不过并定位到具体标题。"""
    plan = plan_for("- [ ] 报告必须包含「市场概览」章节")
    document = Document()
    document.add_heading("市场调研报告", level=0)
    document.add_heading("市场概览", level=1)  # 标题下无正文
    document.save(str(tmp_path / "报告.docx"))

    report = await run_verify(tmp_path, plan, since=0)

    assert report.passed is False and report.layer == "L1"
    empty = next(item for item in report.failures if item.id.startswith("L1-EMPTY"))
    assert "市场概览" in empty.reason and "空章节" in empty.requirement


async def test_l1_url_count_and_datapoint_rules(tmp_path: Path) -> None:
    """L1 引用 URL 数 ≥ N 与数据点判据：先判不过（给出实际值），补齐后通过。"""
    plan = plan_for("- [ ] 引用至少 3 个来源 URL\n- [ ] 数据点：2024 年营收 12.5 亿元\n")
    target = tmp_path / "报告.docx"
    write_docx(target, "报告", ["来源：https://a.example.com/1"], [])

    report = await run_verify(tmp_path, plan, since=0)
    assert report.passed is False and report.layer == "L1"
    urls = next(item for item in report.failures if item.id == "L1-C1")
    assert "实际引用 1 个 URL" in urls.reason and "≥ 3" in urls.reason
    datapoint = next(item for item in report.failures if item.id == "L1-C2")
    assert "2024 年营收 12.5 亿元" in datapoint.reason

    write_docx(
        target,
        "报告",
        [
            "来源：https://a.example.com/1、https://b.example.com/2、https://c.example.com/3",
            "2024 年营收 12.5 亿元，同比增长 8%。",
        ],
        [],
    )
    report = await run_verify(tmp_path, plan, since=0)
    assert report.passed is True
    assert next(item for item in report.items if item.id == "L1-C1").evidence


async def test_l0_scans_only_artifacts_written_in_this_run(tmp_path: Path) -> None:
    """L0 产物范围：只校验本轮写入的交付物，历史旧产物不参与（避免误判）。"""
    plan = plan_for("- [ ] 报告必须包含「市场概览」章节\n")
    old = tmp_path / "历史产物.docx"
    write_docx(old, "历史产物", ["历史产物：TODO 待补充"], [])
    os.utime(old, (time.time() - 7200, time.time() - 7200))
    rate_docx(tmp_path, headings=["市场概览"])

    report = await run_verify(tmp_path, plan, since=time.time() - 60)

    assert report.passed is True
    assert not any("占位符" in item.requirement for item in report.failures)

    # 不限时间下界时，历史旧产物进入校验并判不过（证明过滤确实生效）
    report_all = await run_verify(tmp_path, plan)
    assert report_all.passed is False
    assert any("TODO" in item.reason for item in report_all.failures)


async def test_pptx_and_xlsx_structure_gates(tmp_path: Path) -> None:
    """L0/L1 覆盖 pptx 与 xlsx：真打开 + 关键词命中 + 按单位的行数下限。"""
    from openpyxl import Workbook
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "市场概览"
    slide.placeholders[1].text = "市场规模 12 亿元"
    presentation.save(str(tmp_path / "汇报.pptx"))

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["指标", "数值"])
    for index in range(4):
        sheet.append([f"指标{index}", index])
    workbook.save(str(tmp_path / "数据.xlsx"))

    plan = plan_for("- [ ] 汇报必须包含「市场概览」章节\n- [ ] 数据表至少 3 行\n")
    report = await run_verify(tmp_path, plan, since=0)

    assert report.passed is True
    assert any("python-pptx" in item.evidence for item in report.items)
    assert any("openpyxl" in item.evidence for item in report.items)
    assert report.to_dict()["passed"] is True


# --------------------------------------------------------------------------- #
# UT-M4-3 回灌只修不达标项
# --------------------------------------------------------------------------- #


async def test_ut_m4_3_repair_only_touches_failing_items(tmp_path: Path) -> None:
    """UT-M4-3：回灌指令只含未通过条目；已达标段落回灌后逐字不变（哈希比对）。"""
    work_dir = tmp_path / "ws"
    config = make_config(WORKSPACE_ROOT=str(tmp_path))
    callbacks = FakeCallbacks()
    bridge = OfficeBridge(config, callbacks)
    llm = ScriptedVerifyLLM(work_dir)
    dispatch = make_dispatch(
        output_dir=str(work_dir),
        specification=TaskSpecification(checklist=CHECKLIST),
    )

    with patched_llm(llm):
        await bridge.execute_task(dispatch)

    # ① 只发生一轮回灌，且回灌指令只列未通过条目（已达标条目不出现在输入里）
    assert len(llm.repair_prompts) == 1
    prompt = llm.repair_prompts[0]
    assert "风险提示" in prompt and "L1-C2" in prompt
    assert "市场概览" not in prompt
    assert "期望：" in prompt and "实际：" in prompt
    # 回灌轮带「已达标内容逐字保留」的系统约束
    repair_system = llm.system_prompts[-1]
    assert "已达标内容必须逐字保留" in repair_system

    # ② 已达标段落文本在回灌后逐字不变（回灌前快照哈希 ⊆ 回灌后哈希）
    assert llm.preexisting_digests, "回灌前应已读到首轮交付物"
    final = docx_paragraphs(work_dir / "报告.docx")
    final_digests = {digest(text) for text in final}
    for before in llm.preexisting_digests:
        assert before in final_digests
    for paragraph in PASSING_PARAGRAPHS:
        assert clean_llm_text(paragraph) in final
    assert MISSING_HEADING in final and clean_llm_text(MISSING_PARAGRAPH) in final

    # ③ 通过后走正常交付路径
    assert "completed" in callbacks.kinds()
    assert "verify_blocked" not in callbacks.kinds()
    assert "failed" not in callbacks.kinds()


async def test_repair_instruction_excludes_passing_items(tmp_path: Path) -> None:
    """回灌输入构造：未通过条目带「条目 + 期望 + 实际」，已达标项不得出现。"""
    plan = plan_for(CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"])

    report = await run_verify(tmp_path, plan, since=0)
    instruction = report.repair_instruction(2)

    assert "风险提示" in instruction and "L1-C2" in instruction
    assert "市场概览" not in instruction
    assert "第 2/5 轮" in instruction


# --------------------------------------------------------------------------- #
# UT-M4-4 五轮回灌仍不过 → verify_blocked
# --------------------------------------------------------------------------- #


async def test_ut_m4_4_five_rounds_then_blocked_no_sixth(tmp_path: Path) -> None:
    """UT-M4-4：连续 5 轮不过 → 进 verify_blocked，不再发起第 6 轮。"""
    work_dir = tmp_path / "ws"
    config = make_config(WORKSPACE_ROOT=str(tmp_path))
    callbacks = FakeCallbacks()
    bridge = OfficeBridge(config, callbacks)
    llm = ScriptedVerifyLLM(work_dir, fix_repair=False)  # 模型始终不补齐缺失章节
    dispatch = make_dispatch(
        output_dir=str(work_dir),
        specification=TaskSpecification(checklist=CHECKLIST),
    )

    with patched_llm(llm):
        await bridge.execute_task(dispatch)

    # 初始 1 轮 + 回灌 5 轮 = 6 轮工具循环，每轮 2 次 LLM 调用（写工具 + 收尾）
    assert len(llm.repair_prompts) == MAX_REPAIR_ROUNDS == 5
    assert llm.calls == 12
    assert not any("第 6/" in prompt for prompt in llm.repair_prompts)

    blocked = callbacks.of("verify_blocked")
    assert len(blocked) == 1
    _kind, task_id, session_id, report = blocked[0]
    assert task_id == "task-verify-001"
    assert session_id == "office_task-ver"
    assert report["status"] == "failed"
    assert report["repair_rounds"] == 5
    # 「卡在哪一条」明细：条目 ID + 要求 + 期望 + 实际
    failure = report["must_pass_failures"][0]
    assert failure["id"] == "L1-C2"
    assert "风险提示" in failure["name"]
    assert "风险提示" in failure["expected"]
    assert "缺失" in failure["actual"]
    assert failure["status"] == "failed" and failure["layer"] == "L1"

    # 阻断后不再上报 completed / failed（等用户在对话区决定）
    assert "completed" not in callbacks.kinds()
    assert "failed" not in callbacks.kinds()


async def test_ut_m4_4_loop_stops_after_max_repair_rounds(tmp_path: Path) -> None:
    """UT-M4-4（单元层）：回灌回调恰好被调用 5 次，第 5 轮后不再重试。"""
    plan = plan_for(CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"])
    calls: list[int] = []

    async def repair(instruction: str, round_index: int) -> None:
        calls.append(round_index)

    report = await run_verify_loop(
        tmp_path, plan, llm=None, repair=repair, since=0
    )

    assert calls == [1, 2, 3, 4, 5]
    assert report.passed is False and report.repair_round == 5
    assert report.layer == "L1"


async def test_blocked_payload_matches_hub_contract(tmp_path: Path) -> None:
    """阻断上报 payload 字段与中枢既有契约对齐（task_id/session_id/reason/report）。"""
    plan = plan_for(CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"])
    report = await run_verify(tmp_path, plan, since=0)
    report.repair_round = 5

    payload = report.to_blocked_payload("task-1", "office_task-1")

    assert set(payload) == {"task_id", "session_id", "reason", "report"}
    assert payload["task_id"] == "task-1"
    assert "卡在 L1" in payload["reason"] and "风险提示" in payload["reason"]
    body = payload["report"]
    assert {"status", "source", "must_pass_failures", "checks", "repair_rounds", "summary"} <= set(
        body
    )
    assert body["source"] == "checklist" and body["repair_rounds"] == 5
    assert {"layer", "name", "cmd", "status"} <= set(body["must_pass_failures"][0])


# --------------------------------------------------------------------------- #
# L2 语义闸：必须给证据位置
# --------------------------------------------------------------------------- #


async def test_l2_conclusion_without_evidence_fails(tmp_path: Path) -> None:
    """L2：只给结论不给证据位置 → 判为不通过（ADR-065 M4 硬规则）。"""
    plan = plan_for(SEMANTIC_CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"], paragraphs=["建议：优先拓展华东市场。"])
    judge = JudgeLLM('{"items":[{"id":"C1","passed":true,"evidence":"","reason":""}]}')

    report = await run_verify(tmp_path, plan, llm=judge, since=0)

    assert report.passed is False and report.layer == "L2"
    item = report.failures[0]
    assert item.id == "L2-C1"
    assert "只给结论未给证据位置" in item.reason


async def test_l2_evidence_without_location_fails(tmp_path: Path) -> None:
    """L2：结论带「证据」但没有可核验位置（第几段/第几张幻灯片）→ 不通过。"""
    plan = plan_for(SEMANTIC_CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"], paragraphs=["建议：优先拓展华东市场。"])
    judge = JudgeLLM('{"items":[{"id":"C1","passed":true,"evidence":"文档已满足要求","reason":""}]}')

    report = await run_verify(tmp_path, plan, llm=judge, since=0)

    assert report.passed is False and report.layer == "L2"
    assert "证据未给出可核验的位置" in report.failures[0].reason


async def test_l2_evidence_with_location_passes(tmp_path: Path) -> None:
    """L2：逐条判 + 给出证据位置（段落/幻灯片/引号标题）→ 通过。"""
    plan = plan_for(SEMANTIC_CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"], paragraphs=["建议：优先拓展华东市场。"])
    judge = JudgeLLM(
        '```json\n{"items":[{"id":"C1","passed":true,'
        '"evidence":"第 1 段：建议优先拓展华东市场。","reason":""}]}\n```'
    )

    report = await run_verify(tmp_path, plan, llm=judge, since=0)

    assert report.passed is True
    assert report.layer == "passed"
    passed_item = next(item for item in report.items if item.id == "L2-C1")
    assert "第 1 段" in passed_item.evidence
    assert judge.calls == 1
    assert "证据位置" in judge.prompts[0] and "只给结论不给证据位置" in judge.prompts[0]


async def test_l2_missing_conclusion_and_llm_error_fail(tmp_path: Path) -> None:
    """L2：缺少该条结论 / 判分模型不可用 → 一律按未通过计（不许假成功）。"""
    plan = plan_for(SEMANTIC_CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"])

    empty = await run_verify(tmp_path, plan, llm=JudgeLLM('{"items":[]}'), since=0)
    assert empty.passed is False and "缺少该条结论" in empty.failures[0].reason

    broken = await run_verify(
        tmp_path, plan, llm=JudgeLLM(LLMError("judge down")), since=0
    )
    assert broken.passed is False and "L2 语义判分不可用" in broken.failures[0].reason

    no_llm = await run_verify(tmp_path, plan, llm=None, since=0)
    assert no_llm.passed is False and "未提供判分模型" in no_llm.failures[0].reason


# --------------------------------------------------------------------------- #
# 判据来源与闸门生效范围
# --------------------------------------------------------------------------- #


async def test_plan_criteria_come_only_from_hub_specification() -> None:
    """判据只来自中枢下发的规格/checklist；无判据时不生成任何自编判据。"""
    specification = TaskSpecification(
        spec="产出《季度分析.docx》，正文给出结论。",
        tasks="1. 收集数据\n2. 产出文档",
        checklist=CHECKLIST,
    )
    plan = VerifyPlan.from_specification(specification)

    assert plan.source == "checklist"
    assert [item.id for item in plan.checklist] == ["C1", "C2"]
    assert plan.deliverables == ["季度分析.docx"]
    assert [rule.kind for rule in plan.content_rules] == ["keywords", "keywords"]

    empty = VerifyPlan.from_specification(None)
    assert empty.source == "none" and empty.has_criteria is False
    assert empty.content_rules == [] and empty.semantic == [] and empty.deliverables == []


async def test_gate_skipped_without_criteria_and_artifacts(tmp_path: Path) -> None:
    """无中枢判据且无交付物 → 闸门不生效（不据此判成功，也不误判失败）。"""
    report = await run_verify(tmp_path, VerifyPlan.from_specification(None), since=0)

    assert report.passed is True and report.items == []
    assert "未生效" in report.note


async def test_spec_only_plan_skips_semantic_judging(tmp_path: Path) -> None:
    """无 checklist 时只做 L0/L1（不做 L2 语义逐条判，避免误判规格正文）。"""
    specification = TaskSpecification(
        spec="产出《季度分析.docx》。\n1. 收集数据\n2. 正文给出结论。",
        tasks="1. 收集数据\n2. 产出文档",
        checklist=None,
    )
    plan = VerifyPlan.from_specification(specification)

    assert plan.source == "spec"
    assert plan.semantic == [] and plan.content_rules == []
    assert plan.deliverables == ["季度分析.docx"]

    write_docx(tmp_path / "季度分析.docx", "季度分析报告", ["市场规模 12 亿元。"], ["市场概览"])
    report = await run_verify(tmp_path, plan, llm=JudgeLLM('{"items":[]}'), since=0)
    assert report.passed is True
    assert all(item.layer in ("L0", "L1") for item in report.items)


async def test_tool_loop_repair_round_injects_constraint(tmp_path: Path) -> None:
    """工具循环：回灌轮注入「只修未通过项」系统约束并使用 testing 阶段。"""
    from app.executor.tool_loop import ToolLoopRunner, default_registry

    config = make_config(WORKSPACE_ROOT=str(tmp_path))
    llm = ScriptedVerifyLLM(tmp_path)
    steps: list[tuple[int, str, str]] = []

    async def on_progress(progress: int, stage: str, message: str) -> None:
        steps.append((progress, stage, message))

    runner = ToolLoopRunner(config, llm, default_registry(tmp_path, config), on_progress=on_progress)
    outcome = await runner.run(
        "task-1", "自动交付质量闸未通过：请只修复未通过条目", asyncio.Event(), repair_round=2
    )

    assert outcome.stopped_reason == "final_answer"
    assert "已达标内容必须逐字保留" in llm.system_prompts[0]
    assert "第 2/5 轮" in llm.system_prompts[0]
    assert steps[0][1] == "testing" and "第 2/5 轮" in steps[0][2]

    # 缺省（非回灌轮）不注入修复约束，阶段仍为 planning（既有行为不变）
    llm2 = ScriptedVerifyLLM(tmp_path)
    runner2 = ToolLoopRunner(
        config, llm2, default_registry(tmp_path, config), on_progress=on_progress
    )
    await runner2.run("task-2", "写一份报告", asyncio.Event())
    assert "已达标内容必须逐字保留" not in llm2.system_prompts[0]


@pytest.mark.parametrize("checklist", ["", None])
async def test_empty_checklist_has_no_criteria(checklist: str | None) -> None:
    """空 checklist / 未下发 checklist → 无任务判据（节点不得自编）。"""
    plan = VerifyPlan.from_text(checklist)
    assert plan.has_criteria is False
    assert plan.checklist == [] and plan.content_rules == []


async def test_llm_outage_during_repair_keeps_failure_semantics(tmp_path: Path) -> None:
    """回灌轮 LLM 不可用 → 沿用既有 llm_unavailable 失败语义（不误报内部错误）。"""
    work_dir = tmp_path / "ws"
    config = make_config(WORKSPACE_ROOT=str(tmp_path))
    callbacks = FakeCallbacks()
    bridge = OfficeBridge(config, callbacks)

    class OutageLLM:
        """首轮正常产出缺失章节；进入回灌轮后 LLM 不可用。"""

        model = "fake-model"

        def __init__(self) -> None:
            self.calls = 0

        async def chat(self, messages, tools=None):  # noqa: ANN001, ANN202
            self.calls += 1
            if self.calls == 1:
                return tool_call_response("c1", initial_blocks())
            if self.calls == 2:
                return final_response()
            raise LLMError("provider down")

    dispatch = make_dispatch(
        output_dir=str(work_dir),
        specification=TaskSpecification(checklist=CHECKLIST),
    )
    with patched_llm(OutageLLM()):
        await bridge.execute_task(dispatch)

    failed = callbacks.of("failed")
    assert len(failed) == 1 and "llm_unavailable" in failed[0][2]
    assert "completed" not in callbacks.kinds()


async def test_verify_blocked_uses_existing_event_channel(tmp_path: Path) -> None:
    """阻断上报通道：回调宿主只提供通用事件发送时，仍按既有事件名与字段发出。

    这条覆盖 office 节点现行生效路径（hub_client 未暴露专用 report_verify_blocked），
    断言线上 payload 与中枢 _handle_task_verify_blocked 读取的字段一致。
    """
    sent: list[tuple[str, dict]] = []

    class ClientLike:
        """模拟 HubClient：只暴露 report_failed 与通用 _send_event。"""

        async def report_failed(self, task_id: str, reason: str) -> None:
            """占位（回调宿主解析用）。"""

        async def _send_event(self, event: str, payload: dict) -> None:
            sent.append((event, payload))

    owner = ClientLike()
    callbacks = FakeCallbacks()
    callbacks.on_failed = owner.report_failed  # noqa: SLF001 — 测试注入绑定方法
    callbacks.on_verify_blocked = None  # 模拟 main.py 未接线：走回调宿主的既有通道
    config = make_config(WORKSPACE_ROOT=str(tmp_path))
    bridge = OfficeBridge(config, callbacks)

    plan = plan_for(CHECKLIST)
    rate_docx(tmp_path, headings=["市场概览"])
    report = await run_verify(tmp_path, plan, since=0)
    report.repair_round = 5

    await bridge._report_verify_blocked("task-9", report)  # noqa: SLF001

    assert len(sent) == 1
    event, payload = sent[0]
    assert event == "agent.task_verify_blocked"
    assert set(payload) == {"task_id", "session_id", "reason", "report"}
    assert payload["task_id"] == "task-9" and payload["session_id"] == "office_task-9"
    assert "卡在 L1" in payload["reason"]
    body = payload["report"]
    assert body["status"] == "failed" and body["repair_rounds"] == 5
    assert body["must_pass_failures"][0]["id"] == "L1-C2"
    assert "风险提示" in body["must_pass_failures"][0]["name"]
