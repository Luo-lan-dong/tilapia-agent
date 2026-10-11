"""素材读取工具单元测试（ADR-065 M1：UT-M1-1 / UT-M1-2 / UT-M1-3）。"""

from pathlib import Path

import pytest

from app.tools.base import ToolError
from app.tools.read_tool import (
    MAX_OUTPUT_CHARS,
    ReadCsvTool,
    ReadDocxTool,
    ReadTextTool,
    ReadXlsxTool,
)


def _make_docx(path: Path) -> None:
    """构造含标题/段落/表格的 docx 素材。"""
    from docx import Document

    document = Document()
    document.add_heading("季度总结", level=1)
    document.add_heading("背景", level=2)
    document.add_paragraph("这是一段正文素材。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "CPU"
    table.cell(1, 1).text = "8核"
    document.save(str(path))


def _make_xlsx(path: Path) -> None:
    """构造多 sheet + 空单元格的 xlsx 素材。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["名称", "数量", "备注"])
    sheet.append(["甲", None, "无备注"])
    sheet.append(["乙", 3, None])
    workbook.create_sheet("空表")
    workbook.save(str(path))


class TestReadDocx:
    """UT-M1-1：docx 标题层级 / 段落 / 表格均被抽取。"""

    def test_extracts_heading_paragraph_table(self, tmp_path: Path) -> None:
        _make_docx(tmp_path / "素材.docx")
        result = ReadDocxTool(tmp_path).execute({"filename": "素材.docx"})
        assert "# 季度总结" in result
        assert "## 背景" in result
        assert "这是一段正文素材。" in result
        assert "[[TABLE 1]]" in result
        assert "| 指标 | 数值 |" in result
        assert "| CPU | 8核 |" in result

    def test_reads_nested_relative_path(self, tmp_path: Path) -> None:
        (tmp_path / "素材").mkdir()
        _make_docx(tmp_path / "素材" / "需求.docx")
        result = ReadDocxTool(tmp_path).execute({"filename": "素材/需求.docx"})
        assert "# 季度总结" in result
        assert "1 段落 / 2 标题 / 1 表格" in result

    def test_rejects_non_docx_format(self, tmp_path: Path) -> None:
        (tmp_path / "note.txt").write_text("纯文本", encoding="utf-8")
        with pytest.raises(ToolError) as excinfo:
            ReadDocxTool(tmp_path).execute({"filename": "note.txt"})
        assert "仅支持" in str(excinfo.value)


class TestReadXlsx:
    """UT-M1-2：多 sheet 逐表返回，空单元格不破坏列对齐。"""

    def test_multi_sheet_with_empty_cells(self, tmp_path: Path) -> None:
        _make_xlsx(tmp_path / "销量.xlsx")
        result = ReadXlsxTool(tmp_path).execute({"filename": "销量.xlsx"})
        assert "[[SHEET 数据]]" in result
        assert "[[SHEET 空表]]" in result
        # 列对齐：所有数据行与表头列数一致（空单元格保留空列位）
        data_lines = [line for line in result.splitlines() if line.startswith("|")][:4]
        assert [line.count("|") for line in data_lines] == [4, 4, 4, 4]
        assert "| 甲 |  | 无备注 |" in result
        assert "| 乙 | 3 |  |" in result
        assert "| 名称 | 数量 | 备注 |" in result

    def test_sheet_filter(self, tmp_path: Path) -> None:
        _make_xlsx(tmp_path / "销量.xlsx")
        result = ReadXlsxTool(tmp_path).execute({"filename": "销量.xlsx", "sheet": "数据"})
        assert "[[SHEET 数据]]" in result
        assert "[[SHEET 空表]]" not in result

    def test_unknown_sheet_raises(self, tmp_path: Path) -> None:
        _make_xlsx(tmp_path / "销量.xlsx")
        with pytest.raises(ToolError) as excinfo:
            ReadXlsxTool(tmp_path).execute({"filename": "销量.xlsx", "sheet": "不存在"})
        assert "工作表不存在" in str(excinfo.value)


class TestReadCsv:
    """CSV：表头 + 前 N 行，编码容错。"""

    def test_header_and_row_limit(self, tmp_path: Path) -> None:
        rows = ["名称,数量", *[f"项目{i},{i}" for i in range(1, 11)]]
        (tmp_path / "清单.csv").write_text("\n".join(rows), encoding="utf-8")
        result = ReadCsvTool(tmp_path).execute({"filename": "清单.csv", "max_rows": 3})
        assert "| 名称 | 数量 |" in result
        assert "| 项目3 | 3 |" in result
        assert "项目4" not in result
        assert "共 10 行数据" in result
        assert "仅显示前 3 行数据" in result

    def test_gbk_fallback(self, tmp_path: Path) -> None:
        (tmp_path / "gbk.csv").write_bytes("名称,数量\n甲,1\n".encode("gbk"))
        result = ReadCsvTool(tmp_path).execute({"filename": "gbk.csv"})
        assert "gbk" in result
        assert "甲" in result


class TestReadText:
    """纯文本：原样返回、超限截断并标注。"""

    def test_reads_markdown(self, tmp_path: Path) -> None:
        (tmp_path / "背景.md").write_text("# 背景\n素材正文", encoding="utf-8")
        result = ReadTextTool(tmp_path).execute({"filename": "背景.md"})
        assert "素材正文" in result
        assert "truncated" not in result

    def test_truncates_over_limit(self, tmp_path: Path) -> None:
        (tmp_path / "长文.txt").write_text("长" * (MAX_OUTPUT_CHARS + 500), encoding="utf-8")
        result = ReadTextTool(tmp_path).execute({"filename": "长文.txt"})
        assert "[[truncated]]" in result
        assert len(result) < MAX_OUTPUT_CHARS + 200

    def test_rejects_unknown_extension(self, tmp_path: Path) -> None:
        (tmp_path / "raw.bin").write_bytes(b"\x00\x01")
        with pytest.raises(ToolError) as excinfo:
            ReadTextTool(tmp_path).execute({"filename": "raw.bin"})
        assert "仅支持" in str(excinfo.value)


class TestReadSafety:
    """UT-M1-3：素材缺失 / 格式不符 / 路径越界 → 可读错误且不崩。"""

    @pytest.mark.parametrize(
        "tool_cls",
        [ReadDocxTool, ReadXlsxTool, ReadCsvTool, ReadTextTool],
    )
    def test_missing_file(self, tmp_path: Path, tool_cls: type) -> None:
        with pytest.raises(ToolError) as excinfo:
            tool_cls(tmp_path).execute({"filename": "不存在.txt"})
        assert "不存在" in str(excinfo.value)

    @pytest.mark.parametrize(
        "tool_cls",
        [ReadDocxTool, ReadXlsxTool, ReadCsvTool, ReadTextTool],
    )
    def test_rejects_path_escape(self, tmp_path: Path, tool_cls: type) -> None:
        outside = tmp_path.parent / "secret.txt"
        outside.write_text("越界素材", encoding="utf-8")
        with pytest.raises(ToolError) as excinfo:
            tool_cls(tmp_path).execute({"filename": "../secret.txt"})
        assert "越出工作目录" in str(excinfo.value)
        with pytest.raises(ToolError) as excinfo:
            tool_cls(tmp_path).execute({"filename": "..\\secret.txt"})
        assert "越出工作目录" in str(excinfo.value)

    def test_rejects_absolute_path(self, tmp_path: Path) -> None:
        target = tmp_path / "abs.txt"
        target.write_text("绝对路径素材", encoding="utf-8")
        for filename in (str(target), "/etc/passwd"):
            with pytest.raises(ToolError) as excinfo:
                ReadTextTool(tmp_path).execute({"filename": filename})
            assert "相对路径" in str(excinfo.value)

    def test_rejects_empty_and_directory(self, tmp_path: Path) -> None:
        (tmp_path / "目录").mkdir()
        with pytest.raises(ToolError):
            ReadTextTool(tmp_path).execute({"filename": "   "})
        with pytest.raises(ToolError) as excinfo:
            ReadTextTool(tmp_path).execute({"filename": "目录"})
        assert "不是文件" in str(excinfo.value)

    def test_wrong_format_returns_readable_error(self, tmp_path: Path) -> None:
        _make_xlsx(tmp_path / "表.xlsx")
        (tmp_path / "文本.txt").write_text("x", encoding="utf-8")
        with pytest.raises(ToolError) as excinfo:
            ReadCsvTool(tmp_path).execute({"filename": "表.xlsx"})
        assert "read_csv 仅支持" in str(excinfo.value)
        # 错误信息不含 traceback 细节且为可读中文原因（工具循环据此回传 LLM）
        assert "Traceback" not in str(excinfo.value)
        with pytest.raises(ToolError):
            ReadXlsxTool(tmp_path).execute({"filename": "文本.txt"})

    def test_corrupt_docx_returns_readable_error(self, tmp_path: Path) -> None:
        (tmp_path / "坏.docx").write_bytes(b"not a real docx")
        with pytest.raises(ToolError) as excinfo:
            ReadDocxTool(tmp_path).execute({"filename": "坏.docx"})
        assert "read_docx 读取失败" in str(excinfo.value)

    def test_invalid_row_param(self, tmp_path: Path) -> None:
        (tmp_path / "清单.csv").write_text("a,b\n1,2", encoding="utf-8")
        for value in (0, -1, "5", True):
            with pytest.raises(ToolError) as excinfo:
                ReadCsvTool(tmp_path).execute({"filename": "清单.csv", "max_rows": value})
            assert "max_rows" in str(excinfo.value)
