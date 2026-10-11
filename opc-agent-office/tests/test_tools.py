"""四件套工具单元测试（Harness U-2/U-3/U-4/U-5）。"""

from pathlib import Path

import pytest

from app.tools.base import ToolError
from app.tools.docx_tool import DocxTool
from app.tools.fetch_tool import WebFetchTool, validate_url
from app.tools.html_text import html_to_markdown
from app.tools.pptx_tool import PptxTool
from app.tools.search_tool import parse_ddg_html


class TestDocxTool:
    """U-2：write_docx 落盘且可被 python-docx 重新打开。"""

    def test_write_docx_roundtrip(self, tmp_path: Path) -> None:
        tool = DocxTool(tmp_path)
        result = tool.execute(
            {
                "filename": "测试报告.docx",
                "title": "调研报告",
                "blocks": [
                    {"type": "heading", "level": 1, "text": "背景"},
                    {"type": "paragraph", "text": "这是第一段。"},
                    {"type": "bullet", "text": "要点一"},
                    {"type": "table", "rows": [["指标", "数值"], ["CPU", "8核"]]},
                ],
            }
        )
        assert "已生成" in result
        target = tmp_path / "测试报告.docx"
        assert target.is_file()

        from docx import Document

        document = Document(str(target))
        texts = [p.text for p in document.paragraphs]
        assert "调研报告" in texts
        assert "这是第一段。" in texts
        assert document.tables[0].cell(0, 0).text == "指标"

    def test_rejects_path_escape(self, tmp_path: Path) -> None:
        tool = DocxTool(tmp_path)
        with pytest.raises(ToolError):
            tool.execute(
                {"filename": "../evil.docx", "title": "t", "blocks": [{"type": "paragraph", "text": "x"}]}
            )


class TestPptxTool:
    """U-3：write_pptx 页数与标题正确。"""

    def test_write_pptx_roundtrip(self, tmp_path: Path) -> None:
        tool = PptxTool(tmp_path)
        result = tool.execute(
            {
                "filename": "汇报.pptx",
                "title": "项目汇报",
                "subtitle": "2026 Q3",
                "slides": [
                    {"title": "背景", "bullets": ["要点1", "要点2"], "notes": "备注"},
                    {"title": "总结", "bullets": []},
                ],
            }
        )
        assert "已生成" in result
        target = tmp_path / "汇报.pptx"
        assert target.is_file()

        from pptx import Presentation

        presentation = Presentation(str(target))
        # 封面 + 2 内容页
        assert len(presentation.slides) == 3
        titles = [slide.shapes.title.text for slide in presentation.slides]
        assert titles == ["项目汇报", "背景", "总结"]


class TestParseDdgHtml:
    """U-4：固定 DDG HTML 样本解析；空结果不抛异常。"""

    SAMPLE = """
    <div class="result results_links results_links_deep web-result">
      <h2 class="result__title">
        <a rel="nofollow" class="result__a" href="//duckduckgo.com/l/?uddg=https%3A%2F%2Fexample.com%2Fa&amp;rut=xyz">
          <b>结果</b>标题一
        </a>
      </h2>
      <a class="result__snippet" href="#">这是<b>摘要</b>内容。</a>
    </div>
    <a class="result__a" href="https://direct.com/b">直接链接</a>
    """

    def test_parse_results(self) -> None:
        results = parse_ddg_html(self.SAMPLE, max_results=5)
        assert len(results) == 2
        assert results[0]["url"] == "https://example.com/a"
        assert "结果标题一" in results[0]["title"]
        assert "摘要" in results[0]["snippet"]
        assert results[1]["url"] == "https://direct.com/b"

    def test_empty_html_returns_empty(self) -> None:
        assert parse_ddg_html("<html><body></body></html>") == []


class TestHtmlToMarkdown:
    """U-5 前置：HTML 正文提取。"""

    def test_basic_conversion(self) -> None:
        html = (
            "<html><head><script>var x=1;</script></head><body>"
            "<h1>标题</h1><p>段落一</p><ul><li>条目A</li><li>条目B</li></ul>"
            "<a href='https://example.com'>链接文字</a>"
            "</body></html>"
        )
        markdown = html_to_markdown(html)
        assert "# 标题" in markdown
        assert "段落一" in markdown
        assert "- 条目A" in markdown
        assert "[链接文字](https://example.com)" in markdown
        assert "var x=1" not in markdown

    def test_truncation(self) -> None:
        html = "<p>" + "长" * 50000 + "</p>"
        markdown = html_to_markdown(html)
        assert "已截断" in markdown


class TestValidateUrl:
    def test_rejects_local(self) -> None:
        for url in ("http://localhost/x", "http://127.0.0.1/x", "ftp://a.com", "不是URL"):
            with pytest.raises(ToolError):
                validate_url(url)

    def test_accepts_http(self) -> None:
        assert validate_url("https://example.com/a") == "https://example.com/a"


class TestWebFetchUpgrade:
    """U-5：Fetcher 失败升级 StealthyFetcher（mock）。"""

    def test_upgrade_to_stealthy(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        # ADR-065 M6：隐身通道默认关闭（OFFICE_ENABLE_STEALTH=false），本用例显式开启
        tool = WebFetchTool(tmp_path, enable_stealth=True)

        class FakePage:
            status = 200
            html_content = "<p>隐身通道正文</p>"

        fetcher_calls: list[str] = []

        import sys
        import types

        scrapling_mod = types.ModuleType("scrapling")
        fetchers_mod = types.ModuleType("scrapling.fetchers")

        def fetcher_get(url, **kwargs):  # noqa: ANN001, ANN202
            fetcher_calls.append("fetcher")
            raise RuntimeError("blocked")

        def stealthy_fetch(url, **kwargs):  # noqa: ANN001, ANN202
            fetcher_calls.append("stealthy")
            return FakePage()

        fetchers_mod.Fetcher = type("F", (), {"get": staticmethod(fetcher_get)})
        fetchers_mod.StealthyFetcher = type("S", (), {"fetch": staticmethod(stealthy_fetch)})
        scrapling_mod.fetchers = fetchers_mod
        monkeypatch.setitem(sys.modules, "scrapling", scrapling_mod)
        monkeypatch.setitem(sys.modules, "scrapling.fetchers", fetchers_mod)

        result = tool.execute({"url": "https://example.com/page"})
        assert fetcher_calls == ["fetcher", "stealthy"]
        assert "隐身通道正文" in result
